"""Private, literal-cell XLSX export of a validated author review report."""

import hashlib
import json
import math
import os
import re
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

try:
    from .private_output import output_directory
    from .source_provenance import validate_binding, validate_provenance, references
except ImportError:
    from private_output import output_directory
    from source_provenance import validate_binding, validate_provenance, references


class ReviewExportError(ValueError):
    """Export guidance that never exposes source text or private paths."""


_INVALID_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')
_MAX_CELL_UNITS = 32767
_MAX_ROWS = 1048576
_SEVERITY_COLORS = {'high': 'FCE0E0', 'medium': 'FFF0C2', 'low': 'E6EEF6'}
_HEADERS = (
    'Finding ID', 'Priority', 'Category', 'Exact raw excerpt', 'Start character',
    'End character (exclusive)', 'Source segment IDs', 'Rationale', 'Author question',
    'Suggested author action', 'Disposition', 'Reviewer',
    'Excerpt encoding',
)


def _text(value, *, long=False):
    if not isinstance(value, str) or (not long and _INVALID_XML.search(value)):
        raise ReviewExportError('Report contains text that Excel cannot represent safely.')
    if not long and len(value.encode('utf-16-le')) // 2 > _MAX_CELL_UNITS:
        raise ReviewExportError('A report field exceeds the Excel cell limit; use the JSON report.')
    return value


def _source_text(value):
    """Keep source literal when XML permits it, otherwise use labelled JSON."""
    # XML parsers normalize literal carriage returns to line feeds. Escape them
    # explicitly rather than losing CRLF identity on workbook round trips.
    if _INVALID_XML.search(value) or '\r' in value:
        return _text(json.dumps(value, ensure_ascii=True)), 'JSON string (escaped)'
    return _text(value), 'Literal raw text'


def _integer(value):
    if type(value) is not int or value < 0 or value > 2**53:
        raise ReviewExportError('Report contains invalid source offsets or coverage counts.')
    return value


def _row(sheet, values):
    if sheet.max_row >= _MAX_ROWS:
        raise ReviewExportError('Report exceeds the Excel row limit; use the JSON report.')
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # openpyxl otherwise interprets leading '=' as a formula. Explicit
            # string typing preserves all prefixes and bytes without apostrophes.
            cell.data_type = 's'
            cell.number_format = '@'
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        cell.font = Font(name='Calibri', size=11, color='172B4D')


def _table(sheet, headers, widths):
    _row(sheet, headers)
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'


def _source_pieces(text):
    # A 5000-character piece remains within Excel's limit even if every
    # character needs a six-character JSON escape for XML-illegal controls.
    for offset in range(0, len(text), 5000):
        yield offset, text[offset:offset + 5000]


def _build(report):
    if not isinstance(report, dict) or report.get('status') not in {
        'complete', 'incomplete', 'failed',
    }:
        raise ReviewExportError('Review report status is missing or invalid.')
    segments = report.get('segments')
    findings = report.get('findings')
    coverage = report.get('coverage')
    if not isinstance(segments, list) or not isinstance(findings, list) or not isinstance(coverage, dict):
        raise ReviewExportError('Review report structure is missing or invalid.')
    chunks = coverage.get('chunks')
    if not isinstance(chunks, list) or type(coverage.get('complete')) is not bool:
        raise ReviewExportError('Review coverage is missing or invalid.')
    total = _integer(coverage.get('total_characters'))
    reviewed = _integer(coverage.get('reviewed_characters'))
    if reviewed > total:
        raise ReviewExportError('Review coverage counts are inconsistent.')
    if report['status'] == 'complete' and (not coverage['complete'] or reviewed != total):
        raise ReviewExportError('A complete report requires complete source coverage.')
    chunk_cursor, actual_reviewed, chunk_ids = 0, 0, set()
    for chunk in chunks:
        index = _integer(chunk['chunk_index'])
        start, end = _integer(chunk['start']), _integer(chunk['end'])
        context_start = _integer(chunk['context_start'])
        context_end = _integer(chunk['context_end'])
        if (index in chunk_ids or start != chunk_cursor
                or not 0 <= context_start <= start < end <= context_end <= total
                or chunk['status'] not in {'complete', 'failed'}):
            raise ReviewExportError('Review chunk coverage is inconsistent.')
        chunk_ids.add(index)
        chunk_cursor = end
        if chunk['status'] == 'complete':
            actual_reviewed += end - start
    expected_status = ('complete' if actual_reviewed == total else
                       'incomplete' if actual_reviewed else 'failed')
    if (chunk_cursor != total or actual_reviewed != reviewed
            or report['status'] != expected_status
            or coverage['complete'] != (expected_status == 'complete')):
        raise ReviewExportError('Report status does not match validated chunk coverage.')

    workbook = Workbook()
    overview = workbook.active
    overview.title = 'Overview'
    _table(overview, ('Author review report', 'Value'), {'A': 32, 'B': 105})
    _row(overview, ('Report status', report['status']))
    overview['B2'].fill = PatternFill('solid', fgColor=(
        'E6EEF6' if report['status'] == 'complete' else 'FCE0E0'))
    _row(overview, ('Status meaning', {
        'complete': ('Source coverage is complete. No findings means no passages were flagged '
                     'by this model run. It does not establish accuracy or publication readiness.'),
        'incomplete': 'Source review is incomplete. An empty findings list is not a clean review.',
        'failed': 'Source review failed. An empty findings list is not a clean review.',
    }[report['status']]))
    _row(overview, ('Finding count', len(findings)))
    for severity in ('high', 'medium', 'low'):
        _row(overview, (f'{severity.title()} priority count',
                        sum(item.get('severity') == severity for item in findings if isinstance(item, dict))))
    _row(overview, ('Coverage complete', str(coverage['complete']).lower()))
    _row(overview, ('Characters reviewed', reviewed))
    _row(overview, ('Total raw characters', total))
    for label, key in (
        ('Raw SHA-256', 'raw_sha256'), ('Model', 'model'),
        ('Prompt version', 'prompt_version'), ('Prompt SHA-256', 'prompt_sha256'),
    ):
        _row(overview, (label, _text(report.get(key))))
    for label, value in (
        ('How to use', 'Read each excerpt in its raw context. Record your disposition and reviewer '
         'in Findings. Verify attribution, consent and disputed facts with the author.'),
        ('Review decisions', 'Disposition and Reviewer are editable author records. Changes in this '
         'workbook are not imported by the CLI and do not automatically approve chapter generation.'),
        ('Human review required', 'Flags are suggestions for nuanced human review, not proof. '
         'Criticism is not automatically false, high priority or unsuitable for publication. '
         'No raw text is changed by this report.'),
        ('High priority', 'Concrete serious allegations, sensitive disclosures or strong reputational risk.'),
        ('Medium priority', 'Contextual criticism, disputed facts or unverified attribution '
         'requiring context or verification.'),
        ('Low priority', 'Personal opinions, ambiguity or wording requiring clarification.'),
        ('Source references', 'Offsets are zero-based Python character positions in the raw transcript. '
         'End offsets are exclusive. Source segments preserves exact text in contiguous pieces.'),
        ('Text encoding', 'Literal raw text is unchanged. Cells marked JSON string (escaped) contain '
         'a lossless JSON representation because Excel XML rejects or normalizes some controls. '
         'Decode that JSON string to recover the exact source. The JSON report remains canonical.'),
        ('Long cells', 'Very long text may exceed Excel row-height display limits. Select the cell '
         'and use the formula bar, or read the canonical JSON report. No stored text is truncated.'),
        ('Draft gate', 'Unresolved high priority findings require human review before chapter generation, '
         'unless an explicit draft override is used with visible warnings.'),
    ):
        _row(overview, (label, value))
    for key in ('schema_version', 'settings_fingerprint'):
        if key in report:
            value = report[key]
            value = _integer(value) if type(value) is int else _text(value)
            _row(overview, (key.replace('_', ' ').capitalize(), value))
    for row in range(2, overview.max_row + 1):
        value = str(overview.cell(row, 2).value)
        overview.row_dimensions[row].height = max(30, math.ceil(len(value) / 90) * 17)

    source_sheet = workbook.create_sheet('Source segments')
    _table(source_sheet, ('Segment ID', 'Piece', 'Start character', 'End character (exclusive)',
                          'Source text', 'Text encoding'),
           {'A': 22, 'B': 10, 'C': 20, 'D': 24, 'E': 110, 'F': 25})
    source_by_id = {}
    source_parts = []
    cursor = 0
    for segment in segments:
        segment_id = _text(segment['segment_id'])
        start, end = _integer(segment['start']), _integer(segment['end'])
        text = _text(segment['text'], long=True)
        if segment_id in source_by_id or not segment_id or start != cursor or end - start != len(text):
            raise ReviewExportError('Source segments have duplicate IDs or inconsistent offsets.')
        source_by_id[segment_id] = (start, end)
        source_parts.append(text)
        cursor = end
        for piece, (local_start, part) in enumerate(_source_pieces(text), 1):
            displayed, encoding = _source_text(part)
            _row(source_sheet, (segment_id, piece, start + local_start,
                                start + local_start + len(part), displayed, encoding))
            source_sheet.row_dimensions[source_sheet.max_row].height = 90
    if cursor != total:
        raise ReviewExportError('Source segments do not cover the reported raw text.')
    source = ''.join(source_parts)
    if (not source.strip() or hashlib.sha256(source.encode('utf-8')).hexdigest()
            != report['raw_sha256']):
        raise ReviewExportError('Report source does not match its raw transcript hash.')

    if 'recording_provenance' in report:
        provenance = report['recording_provenance']
        validate_provenance(source, provenance)
        validate_binding(report, provenance)
        parts_sheet = workbook.create_sheet('Recording parts')
        _table(parts_sheet, ('Order', 'Part ID', 'Local source file', 'Source SHA-256',
                            'Part raw transcript', 'Part raw SHA-256', 'Combined start', 'Combined end'),
               {'A': 12, 'B': 22, 'C': 65, 'D': 68, 'E': 65, 'F': 68, 'G': 20, 'H': 20})
        _row(overview, ('Recording references', 'Recording parts and Recording references map '
             'combined characters to local part text. Separators are metadata, not missing speech. '
             'Text offsets are not audio timestamps; no global timeline is inferred across pauses.'))
        for part in provenance['parts']:
            _row(parts_sheet, (_integer(part['order']), _text(part['id']), _text(part['path']),
                 _text(part['source_sha256']), _text(part['raw_transcript']), _text(part['raw_sha256']),
                 _integer(part['start']), _integer(part['end'])))
        refs_sheet = workbook.create_sheet('Recording references')
        _table(refs_sheet, ('Reference type', 'Reference ID', 'Span kind', 'Part order', 'Part ID',
                           'Combined start', 'Combined end', 'Part local start', 'Part local end', 'Part segment IDs'),
               {'A': 22, 'B': 40, 'C': 20, 'D': 14, 'E': 22, 'F': 20, 'G': 20, 'H': 20, 'I': 20, 'J': 35})
        for kind, items, id_key in [('segment', segments, 'segment_id'),
                                    ('finding', findings, 'finding_id'), ('chunk', chunks, 'chunk_index')]:
            for item in items:
                for ref in references(provenance, item['start'], item['end']):
                    _row(refs_sheet, (_text(kind), _text(str(item[id_key])), _text(ref['kind']), ref.get('order', ''),
                         _text(ref.get('part_id', '')), ref['start'], ref['end'], ref.get('local_start', ''),
                         ref.get('local_end', ''), _text(', '.join(ref.get('local_segment_ids', [])))))
        for sheet in (parts_sheet, refs_sheet):
            sheet.auto_filter.ref = sheet.dimensions

    finding_sheet = workbook.create_sheet('Findings', 1)
    _table(finding_sheet, _HEADERS, {
        'A': 24, 'B': 12, 'C': 24, 'D': 75, 'E': 18, 'F': 24, 'G': 30,
        'H': 55, 'I': 55, 'J': 55, 'K': 20, 'L': 24, 'M': 25,
    })
    ids = set()
    for finding in findings:
        finding_id = _text(finding['finding_id'])
        severity = finding['severity']
        start, end = _integer(finding['start']), _integer(finding['end'])
        excerpt = _text(finding['excerpt'], long=True)
        referenced = finding['segment_ids']
        if (not finding_id or finding_id in ids or severity not in _SEVERITY_COLORS
                or not excerpt or not start < end <= total or source[start:end] != excerpt
                or not isinstance(referenced, list) or not referenced
                or len(set(referenced)) != len(referenced)
                or any(item not in source_by_id for item in referenced)
                or set(referenced) != {key for key, (a, b) in source_by_id.items()
                                       if a < end and b > start}):
            raise ReviewExportError('A finding has invalid identity, priority or exact source references.')
        ids.add(finding_id)
        displayed, encoding = _source_text(excerpt)
        _row(finding_sheet, (
            finding_id, severity, _text(finding['category']), displayed, start, end,
            _text(', '.join(referenced)), _text(finding['rationale']),
            _text(finding.get('author_question', '')), _text(finding['author_action']),
            _text(finding.get('disposition', 'Unreviewed')), _text(finding.get('reviewer', '')),
            encoding,
        ))
        row = finding_sheet.max_row
        finding_sheet.row_dimensions[row].height = 105
        finding_sheet.cell(row, 2).fill = PatternFill('solid', fgColor=_SEVERITY_COLORS[severity])
        for column in (11, 12):
            finding_sheet.cell(row, column).fill = PatternFill('solid', fgColor='E8F4FC')
    disposition = DataValidation(
        type='list', formula1='"Unreviewed,Clarify,Verify,Revise,Retain,Exclude"', allow_blank=True,
    )
    disposition.errorTitle = 'Choose a disposition'
    disposition.error = 'Select a disposition or leave this field blank.'
    disposition.showErrorMessage = True
    finding_sheet.add_data_validation(disposition)
    disposition.add(f'K2:K{max(2, finding_sheet.max_row)}')

    coverage_sheet = workbook.create_sheet('Coverage')
    _table(coverage_sheet, ('Chunk', 'Start character', 'End character (exclusive)',
                            'Context start', 'Context end', 'Status', 'Safe error code'),
           {'A': 12, 'B': 20, 'C': 24, 'D': 20, 'E': 20, 'F': 18, 'G': 50})
    for chunk in chunks:
        _row(coverage_sheet, (
            _integer(chunk['chunk_index']), _integer(chunk['start']), _integer(chunk['end']),
            _integer(chunk['context_start']), _integer(chunk['context_end']),
            _text(chunk['status']), _text(chunk.get('error', '')),
        ))
        coverage_sheet.row_dimensions[coverage_sheet.max_row].height = 36
    for sheet in (finding_sheet, source_sheet, coverage_sheet):
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_options.horizontalCentered = True
    workbook.properties.creator = 'Voice Transcription Engine'
    workbook.properties.title = 'Author review report'
    workbook.properties.description = 'Human review suggestions with exact raw source references.'
    return workbook


def export_review(report: dict, path: Path) -> None:
    """Export a complete/incomplete/failed report atomically, without overwriting.

    All source and model text is stored as literal Excel strings. Source segment
    pieces preserve complete text; other oversized fields fail safely rather than
    relying on openpyxl's automatic 32767-character truncation.
    """
    temporary = None
    try:
        path = Path(path)
        if path.suffix.lower() != '.xlsx':
            raise ReviewExportError('Review workbook output must use the .xlsx extension.')
        if os.path.lexists(path):
            raise ReviewExportError('Review workbook already exists; no file was overwritten.')
        workbook = _build(report)
        directory = output_directory(path.parent)
        path = directory / path.name
        fd, temporary = tempfile.mkstemp(prefix='.review-', dir=directory)
        with os.fdopen(fd, 'w+b') as stream:
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        os.link(temporary, path)
    except ReviewExportError:
        raise
    except FileExistsError:
        raise ReviewExportError('Review workbook already exists; no file was overwritten.') from None
    except (KeyError, TypeError, AttributeError, ValueError):
        raise ReviewExportError('Review report is invalid or cannot be represented safely in Excel.') from None
    except Exception:
        raise ReviewExportError('Cannot save review workbook; check output access and free space.') from None
    finally:
        if temporary is not None:
            try:
                os.unlink(temporary)
            except OSError:
                pass
