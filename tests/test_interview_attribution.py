"""Synthetic turn/mapping, independent-family, resume and privacy contracts."""

import hashlib
import json
import shutil
import os
import wave
from types import SimpleNamespace
from unittest.mock import MagicMock

import httpx2
from openai import OpenAI
import pytest

from src.author_workflow import AuthorOptions
from src.chapters import ChapterOptions
from src.cli import main, entrypoint
from src.interview_attribution import (
    AttributedInterview, InterviewOptions, DIARIZATION_MODEL, diarization_configuration,
    diarize, input_record,
)
from src.model_config import TranscriptionOptions, EditingOptions, _fingerprint
from src.ordered_interview import OrderedInterview
from src.pipeline import Pipeline, PipelineError
from src.source_provenance import validate_provenance, validate_binding, validate_chapter_binding
from src.transcriber import Transcriber, TranscriptionError
from openpyxl import load_workbook


@pytest.fixture
def interview_provider():
    client = MagicMock()
    client.high = False
    def audio(**parameters):
        if parameters['model'] != DIARIZATION_MODEL:
            return SimpleNamespace(text='Original raw transcript.\r\nAll original words.')
        return {'text': 'A: A question?\nB: An answer.\nC: Extra words.', 'segments': [
            {'speaker': 'A', 'start': 0.0, 'end': 0.4, 'text': 'A question?'},
            {'speaker': 'B', 'start': 0.3, 'end': 0.8, 'text': 'An answer.'},
            {'speaker': 'C', 'start': 0.8, 'end': 1.0, 'text': 'Extra words.'}]}
    client.audio.transcriptions.create.side_effect = audio
    def chat(**parameters):
        supplied = json.loads(parameters['messages'][-1]['content'])
        name = parameters['response_format']['json_schema']['name']
        if name == 'faithful_transcript_edit':
            body = dict(chunk_index=supplied['chunk_index'], text=supplied['text'], speaker_uncertain=False)
        elif name == 'faithful_turn_group_edit':
            body = dict(group_index=supplied['group_index'], edits=[
                {**turn, 'speaker_uncertain': False} for turn in supplied['turns']])
        elif name == 'source_grounded_author_review':
            findings = []
            if client.high:
                findings = [dict(category='allegation', severity='high', reason_code='serious_allegation',
                                 excerpt=supplied['text'], start=0, end=len(supplied['text']))]
            body = dict(chunk_index=supplied['chunk_index'], fully_reviewed=True,
                        reviewed_start=supplied['core_start'], reviewed_end=supplied['core_end'], findings=findings)
        else:
            body = dict(chunk_index=supplied['chunk_index'], passages=[
                dict(unit_ids=[unit['unit_id']], text=unit['text'].strip(), kind='verbatim_excerpt')
                for unit in supplied['source_units']], coverage_omissions=[])
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',
            message=SimpleNamespace(content=json.dumps(body), refusal=None))])
    client.chat.completions.create.side_effect = chat
    return client


def state(job):
    return json.loads((job / 'manifest.json').read_text())


def artifact(job, stage, filename):
    return next(job / name for name in state(job)['stages'][stage]['artifacts'] if name.endswith('/' + filename))


def args(source, output, *extra):
    return ['--workflow', '--interview', '--input', str(source), '--output-folder', str(output),
            '--interviewer-name', 'Synthetic Interviewer', '--interviewee-name', 'Synthetic Guest', *extra]


def test_additive_all_stages_and_resume(monkeypatch, synthetic_media, tmp_path, interview_provider, capsys):
    source = synthetic_media()
    output = tmp_path / 'outputs'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(source, output, '--chapters', 'both', '--speaker-map', '1:1:A=interviewer')
    assert entrypoint(argv) == 0
    original = next(path for path in output.iterdir() if (path / 'transcription.txt').exists())
    family = next(path.parent for path in (output / 'attributed').rglob('transcription.txt'))
    original_bytes = {path.relative_to(original): path.read_bytes() for path in original.rglob('*') if path.is_file()}
    assert (original / 'transcription.txt').read_bytes() == b'Original raw transcript.\r\nAll original words.'
    raw = (family / 'transcription.txt').read_text()
    assert 'Synthetic Interviewer' in raw and 'Speaker B' in raw and 'Speaker C' in raw
    assert 'Synthetic Guest' not in raw
    provenance = json.loads((family / 'provenance.json').read_text())
    validate_provenance(raw, provenance)
    assert_part_artifacts(output, family, provenance)
    turns = provenance['attribution']['turns']
    assert turns[0]['role'] == 'interviewer' and turns[1]['role'] is None
    assert turns[0]['overlap_detected'] and turns[1]['overlap_detected']
    report = json.loads(artifact(family, 'author_review', 'review_report.json').read_text())
    validate_binding(report, provenance)
    book = load_workbook(artifact(family, 'author_review', 'review_report.xlsx'))
    assert book['Speaker turns'].cell(2, 7).value == 'user_confirmed_mapping'
    assert book['Recording parts'].cell(2, 5).value == provenance['parts'][0]['raw_transcript']
    chapters = json.loads(artifact(family, 'chapters', 'chapter_drafts.json').read_text())
    validate_chapter_binding(chapters, provenance)
    for style in ('interview', 'narrative'):
        assert 'identity evidence: unidentified' in artifact(family, 'chapters', f'chapter_{style}.txt').read_text()
    chapter_text = artifact(family, 'chapters', 'chapter_interview.txt').read_text()
    assert 'Names and roles appear only for user-confirmed mappings' in chapter_text
    assert 'Speaker roles are unassigned.' not in chapter_text
    assert 'Speaker roles are unassigned.' in artifact(original, 'chapters', 'chapter_interview.txt').read_text()
    edited = (family / 'derivative_readability.txt').read_text()
    assert 'unidentified' in edited and 'A question?' in edited
    calls = interview_provider.audio.transcriptions.create.call_count, interview_provider.chat.completions.create.call_count
    assert entrypoint([*argv, '--resume']) == 0
    assert calls == (interview_provider.audio.transcriptions.create.call_count, interview_provider.chat.completions.create.call_count)
    assert original_bytes == {path.relative_to(original): path.read_bytes() for path in original.rglob('*') if path.is_file()}
    logs = capsys.readouterr()
    assert any('diarization' in line for line in logs.out.splitlines())
    for private in ('Synthetic Interviewer', 'Synthetic Guest', 'A question?', str(tmp_path), source.name):
        assert private not in logs.out + logs.err
    for call in interview_provider.audio.transcriptions.create.call_args_list:
        if call.kwargs['model'] == DIARIZATION_MODEL:
            assert set(call.kwargs) == {'file', 'model', 'response_format', 'chunking_strategy'}
    if os.name == 'posix':
        for path in family.rglob('*'):
            assert path.stat().st_mode & 0o077 == 0


@pytest.mark.parametrize('change', ['raw', 'part', 'cache', 'polish', 'review', 'xlsx', 'chapter'])
def test_tamper_independent_resume(monkeypatch, synthetic_media, tmp_path, interview_provider, change):
    source, output = synthetic_media(), tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(source, output, '--chapters', 'both')
    assert main(argv) == 0
    family = next(path.parent for path in (output / 'attributed').rglob('transcription.txt'))
    targets = {'raw': family / 'transcription.txt', 'part': family / 'part-000001_transcription.txt', 'cache': next((output / 'attributed').rglob('provider_responses.json')),
               'polish': family / 'derivative_readability.txt', 'review': artifact(family, 'author_review', 'review_report.json'),
               'xlsx': artifact(family, 'author_review', 'review_report.xlsx'), 'chapter': artifact(family, 'chapters', 'chapter_interview.txt')}
    target = targets[change]
    target.write_bytes(target.read_bytes() + b'changed')
    before = target.read_bytes()
    calls = interview_provider.audio.transcriptions.create.call_count, interview_provider.chat.completions.create.call_count
    assert main([*argv, '--resume']) == 1
    assert before == target.read_bytes()
    assert calls == (interview_provider.audio.transcriptions.create.call_count, interview_provider.chat.completions.create.call_count)


def test_mapping_change_reuses_diarization_without_touching_original(monkeypatch, synthetic_media, tmp_path, interview_provider):
    source, output = synthetic_media(), tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(source, output, '--stages', 'raw')
    assert main(argv) == 0
    original = next(path for path in output.iterdir() if (path / 'transcription.txt').exists())
    before = (original / 'manifest.json').read_bytes()
    assert main([*argv, '--resume', '--speaker-map', '1:1:B=interviewee']) == 0
    assert interview_provider.audio.transcriptions.create.call_count == 2
    assert len(list((output / 'attributed').rglob('transcription.txt'))) == 2
    assert (original / 'manifest.json').read_bytes() == before
    assert main([*argv, '--resume', '--speaker-map', '1:1:Z=interviewer']) == 1
    assert interview_provider.audio.transcriptions.create.call_count == 2


@pytest.mark.parametrize('extra', [[], ['--interviewer-name', ''], ['--interview-model', 'whisper-1'],
    ['--speaker-map', 'A=interviewer'], ['--speaker-map', '1:1:A=interviewer', '--speaker-map', '1:1:A=interviewee'],
    ['--language', 'en', '--language', 'fr']])
def test_invalid_before_provider(monkeypatch, tmp_path, extra, capsys):
    ctor = MagicMock(side_effect=AssertionError('No provider initialization'))
    monkeypatch.setattr('src.cli.Transcriber', ctor)
    base = ['--workflow', '--interview', '--input', str(tmp_path / 'private.wav')]
    if extra:
        base += ['--interviewer-name', 'SYNTHETIC_PRIVATE_HOST', '--interviewee-name', 'SYNTHETIC_PRIVATE_GUEST']
    assert main([*base, *extra]) == 1
    ctor.assert_not_called()
    logs = capsys.readouterr()
    assert 'SYNTHETIC_PRIVATE_' not in logs.out + logs.err


@pytest.mark.parametrize('argv', [ ['--interview'], ['--interviewer-name', 'PRIVATE'],
    ['--workflow', '--extract-only', '--interview'], ['--interview-model', 'PRIVATE']])
def test_flags_require_workflow_and_mode(argv, capsys):
    with pytest.raises(SystemExit) as error:
        main(['--input', 'unused.wav', *argv])
    assert error.value.code == 2
    assert 'PRIVATE' not in capsys.readouterr().err


def test_no_mode_has_no_extra_requests(monkeypatch, synthetic_media, tmp_path, interview_provider):
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    assert main(['--workflow', '--input', str(synthetic_media()), '--output-folder', str(tmp_path / 'out'), '--stages', 'raw']) == 0
    assert interview_provider.audio.transcriptions.create.call_count == 1
    assert not (tmp_path / 'out' / 'attributed').exists()


def test_chunk_identity_reset_tail_and_sdk_contract(tmp_path):
    audio = tmp_path / 'synthetic.wav'
    with wave.open(str(audio), 'wb') as wav:
        wav.setparams((1, 2, 16000, 0, 'NONE', 'not compressed'))
        wav.writeframes(b'\0\0' * 36000)
    requests = []
    def respond(request):
        requests.append(request)
        return httpx2.Response(200, json={'text': 'A: Exact tail.', 'segments': [
            {'speaker': 'A', 'start': 0, 'end': 0.1, 'text': 'Exact tail.'}]})
    with httpx2.Client(transport=httpx2.MockTransport(respond)) as transport:
        with OpenAI(api_key='synthetic-placeholder', http_client=transport, max_retries=0) as client:
            transcriber = Transcriber(client=client, options=TranscriptionOptions(chunk_seconds=1))
            payload = diarize(transcriber, audio, diarization_configuration(transcriber))
    assert [r['offset_seconds'] for r in payload['requests']] == [0, 1, 2]
    assert [r['duration_seconds'] for r in payload['requests']] == [1, 1, 0.25]
    for request in requests:
        assert b'filename="audio.wav"' in request.content
        assert b'diarized_json' in request.content and b'chunking_strategy' in request.content
        assert b'prompt' not in request.content and audio.name.encode() not in request.content
    inputs = [{'order': 1, 'source_sha256': 'a', 'audio_sha256': 'b', 'original_raw_sha256': 'c'}]
    family = AttributedInterview(tmp_path / 'out', inputs,
        InterviewOptions('Host', 'Guest', ('1:1:A=interviewer',)), resume=False, enhance=False, author_options=None)
    _, provenance, _ = family._assemble([(inputs[0], payload, 'opaque')])
    turns = provenance['attribution']['turns']
    assert [t['speaker_key'] for t in turns] == ['1:1:A', '1:2:A', '1:3:A']
    assert [t['identity_evidence'] for t in turns] == ['user_confirmed_mapping', 'unidentified', 'unidentified']
    assert [t['audio_start'] for t in turns] == [0, 1, 2]


def test_part_identity_reset_and_missing_speaker(tmp_path):
    parts = [dict(order=i, source_sha256='a', audio_sha256='b', original_raw_sha256='c') for i in (1, 2)]
    payload = {'requests': [{'request': 1, 'offset_seconds': 0, 'duration_seconds': 1,
        'response': {'text': 'Exact words.', 'segments': [dict(speaker='A', start=0, end=1, text='Exact words.')]}}]}
    family = AttributedInterview(tmp_path / 'out', parts, InterviewOptions('Host', 'Guest', ('1:1:A=interviewer',)),
                                 resume=False, enhance=False, author_options=None)
    raw, provenance, _ = family._assemble([(p, payload, 'binding') for p in parts])
    validate_provenance(raw, provenance)
    assert [t['role'] for t in provenance['attribution']['turns']] == ['interviewer', None]
    payload['requests'][0]['response']['segments'][0]['speaker'] = None
    family.options = InterviewOptions('Host', 'Guest')
    raw, _, _ = family._assemble([(parts[0], payload, 'binding')])
    assert 'Speaker unknown' in raw


def test_disagreement_keeps_original_and_lossless_provider_cache(monkeypatch, synthetic_media, tmp_path, interview_provider):
    original_audio = interview_provider.audio.transcriptions.create.side_effect
    def inconsistent(**kw):
        response = original_audio(**kw)
        if isinstance(response, dict):
            response['text'] += ' Words without a turn.'
        return response
    interview_provider.audio.transcriptions.create.side_effect = inconsistent
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    output = tmp_path / 'out'
    assert main(args(synthetic_media(), output, '--stages', 'raw')) == 1
    assert len(list(output.rglob('transcription.txt'))) == 1
    payload = json.loads(next(output.rglob('provider_responses.json')).read_text())
    assert 'Words without a turn.' in payload['requests'][0]['response']['text']


def test_both_families_keep_high_gate(synthetic_media, tmp_path, interview_provider):
    interview_provider.high = True
    options = AuthorOptions(chapter_options=ChapterOptions())
    transcriber = Transcriber(client=interview_provider)
    output = tmp_path / 'out'
    with pytest.raises(PipelineError):
        Pipeline(output, author_options=options, interview_options=InterviewOptions('Host', 'Guest')).process(
            synthetic_media(), transcriber=transcriber)
    assert len(list(output.rglob('review_report.json'))) == 2
    assert not list(output.rglob('chapter_drafts.json'))
    assert len(list(output.rglob('transcription.txt'))) == 2


@pytest.mark.parametrize('bad', ['negative', 'nan', 'missing', 'provider_error'])
def test_bad_response_no_attributed_raw(monkeypatch, synthetic_media, tmp_path, interview_provider, bad):
    original_audio = interview_provider.audio.transcriptions.create.side_effect
    def malformed(**kw):
        response = original_audio(**kw)
        if isinstance(response, dict):
            if bad == 'provider_error':
                raise RuntimeError('SYNTHETIC_PRIVATE_PROVIDER_PAYLOAD')
            if bad == 'negative':
                response['segments'][0]['start'] = -1
            elif bad == 'nan':
                response['segments'][0]['end'] = float('nan')
            else:
                del response['segments']
        return response
    interview_provider.audio.transcriptions.create.side_effect = malformed
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    output = tmp_path / 'out'
    assert main(args(synthetic_media(), output, '--stages', 'raw')) == 1
    assert len(list(output.rglob('transcription.txt'))) == 1
    assert not list((output / 'attributed').rglob('provider_responses.json'))


def test_symlink_collision(synthetic_media, tmp_path, interview_provider):
    source, out = synthetic_media(), tmp_path / 'out'
    transcriber = Transcriber(client=interview_provider)
    identity, _ = Pipeline(out).process(source, transcriber=transcriber)
    job = out / identity
    family = AttributedInterview(out, [input_record(1, job, state(job))], InterviewOptions('Host', 'Guest'),
                                 resume=True, enhance=False, author_options=None)
    family.job.parent.mkdir()
    family.job.symlink_to(tmp_path, target_is_directory=True)
    with pytest.raises(TranscriptionError):
        family.process(transcriber)
    assert interview_provider.audio.transcriptions.create.call_count == 1


def test_ordered_cli_families_keep_part_provenance_and_stage_selection(monkeypatch, synthetic_media, tmp_path, interview_provider):
    sources = [synthetic_media('part-one.wav'), synthetic_media('part-two.wav')]
    manifest = tmp_path / 'interview.json'
    manifest.write_text(json.dumps({'version': 1, 'interview_id': 'synthetic-session',
        'parts': [{'id': f'part-{i}', 'path': s.name} for i, s in enumerate(sources, 1)]}))
    output = tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = ['--workflow', '--interview', '--interview-manifest', str(manifest), '--output-folder', str(output),
            '--interviewer-name', 'Host', '--interviewee-name', 'Guest', '--stages', 'raw,review',
            '--speaker-map', '1:1:A=interviewer', '--speaker-map', '2:1:B=interviewee']
    assert main(argv) == 0
    family = next(path.parent for path in (output / 'attributed').rglob('transcription.txt'))
    provenance = json.loads((family / 'provenance.json').read_text())
    validate_provenance((family / 'transcription.txt').read_text(), provenance)
    assert_part_artifacts(output, family, provenance)
    book = load_workbook(artifact(family, 'author_review', 'review_report.xlsx'))
    assert [book['Recording parts'].cell(row, 5).value for row in (2, 3)] == [
        part['raw_transcript'] for part in provenance['parts']]
    roles = [t['role'] for t in provenance['attribution']['turns']]
    assert roles == ['interviewer', None, None, None, 'interviewee', None]
    assert len(list(output.rglob('review_report.json'))) == 2
    assert not list(output.rglob('derivative_readability.txt'))
    assert not list(output.rglob('chapter_drafts.json'))
    calls = interview_provider.audio.transcriptions.create.call_count, interview_provider.chat.completions.create.call_count
    assert main([*argv, '--resume']) == 0
    assert calls == (interview_provider.audio.transcriptions.create.call_count, interview_provider.chat.completions.create.call_count)


def test_missing_cache_is_a_conflict_before_new_calls(monkeypatch, synthetic_media, tmp_path, interview_provider):
    source, output = synthetic_media(), tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(source, output, '--stages', 'raw')
    assert main(argv) == 0
    import shutil
    shutil.rmtree(next(output.rglob('provider_responses.json')).parent)
    assert main([*argv, '--resume']) == 1
    assert interview_provider.audio.transcriptions.create.call_count == 2


def test_changed_review_configuration_preserves_and_reselects_versions(monkeypatch, synthetic_media, tmp_path, interview_provider):
    source, output = synthetic_media(), tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(source, output, '--stages', 'raw,review')
    assert main(argv) == 0
    calls = interview_provider.chat.completions.create.call_count
    audio_calls = interview_provider.audio.transcriptions.create.call_count
    preserved = {path: path.read_bytes() for name in ('transcription.txt', 'review_report.json', 'review_report.xlsx')
                 for path in output.rglob(name)}
    assert main([*argv, '--resume', '--author-model', 'gpt-6.1-sol']) == 0
    assert interview_provider.chat.completions.create.call_count == calls + 2
    assert interview_provider.audio.transcriptions.create.call_count == audio_calls
    assert all(path.read_bytes() == body for path, body in preserved.items())
    assert main([*argv, '--resume']) == 0
    assert interview_provider.chat.completions.create.call_count == calls + 2


def test_partial_request_failure_preserves_original_and_no_complete_cache(monkeypatch, tmp_path, interview_provider):
    source = tmp_path / 'synthetic.wav'
    with wave.open(str(source), 'wb') as wav:
        wav.setparams((1, 2, 16000, 0, 'NONE', 'not compressed'))
        wav.writeframes(b'\0\0' * 48000)
    audio = interview_provider.audio.transcriptions.create.side_effect
    count = 0
    def fail(**kw):
        nonlocal count
        if kw['model'] == DIARIZATION_MODEL:
            count += 1
            if count == 2:
                raise RuntimeError('PRIVATE_PAYLOAD')
        return audio(**kw)
    interview_provider.audio.transcriptions.create.side_effect = fail
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    output = tmp_path / 'out'
    argv = args(source, output, '--stages', 'raw', '--audio-chunk-seconds', '1')
    assert main(argv) == 1
    assert len(list(output.rglob('transcription.txt'))) == 1
    assert not list(output.rglob('provider_responses.json'))
    interview_provider.audio.transcriptions.create.side_effect = audio
    assert main([*argv, '--resume']) == 0
    assert len(list(output.rglob('transcription.txt'))) == 2


def test_explicit_high_override_applies_independently(monkeypatch, synthetic_media, tmp_path, interview_provider):
    source, output = synthetic_media(), tmp_path / 'out'
    interview_provider.high = True
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(source, output, '--chapters', 'both')
    assert main(argv) == 1
    reviews = {p: p.read_bytes() for p in output.rglob('review_report.json')}
    raws = {p: p.read_bytes() for p in output.rglob('transcription.txt')}
    assert len(reviews) == 2 and not list(output.rglob('chapter_drafts.json'))
    assert main([*argv, '--resume', '--draft-with-unresolved-high']) == 0
    assert {p: p.read_bytes() for p in reviews} == reviews
    assert {p: p.read_bytes() for p in raws} == raws
    chapters = list(output.rglob('chapter_drafts.json'))
    assert len(chapters) == 2
    assert all(json.loads(p.read_text())['allow_unresolved_high'] for p in chapters)


@pytest.mark.parametrize('name', ['Host\nGuest', 'Host\tGuest', 'Host\u202eGuest', 'Host|Guest', 'Host[Guest]', 'Host\u2028Guest'])
def test_names_cannot_spoof_label_metadata(name):
    from src.model_config import ModelConfigurationError
    with pytest.raises(ModelConfigurationError):
        InterviewOptions(name, 'Guest')


def test_fractional_chunk_resume(tmp_path, interview_provider):
    source = tmp_path / 'synthetic.wav'
    with wave.open(str(source), 'wb') as audio:
        audio.setparams((1, 2, 16000, 0, 'NONE', 'not compressed'))
        audio.writeframes(b'\0\0' * 128000)
    output = tmp_path / 'out'
    options = TranscriptionOptions(chunk_seconds=1.123)
    transcriber = Transcriber(client=interview_provider, options=options)
    identity, _ = Pipeline(output, options=options).process(source, transcriber=transcriber)
    job = output / identity
    part = input_record(1, job, state(job))
    config = diarization_configuration(transcriber)
    family = AttributedInterview(output, [part], InterviewOptions('Host', 'Guest'),
                                 resume=False, enhance=False, author_options=None)
    # A simple response fits the fractional final request as well.
    interview_provider.audio.transcriptions.create.side_effect = lambda **kw: {
        'text': 'Words.', 'segments': [dict(speaker='A', start=0, end=0.01, text='Words.')]}
    payload, _ = family._cache(part, transcriber, config)
    family.resume = True
    reloaded, _ = family._cache(part, transcriber, config)
    assert reloaded == payload


def test_identical_content_files_get_distinct_output_families(monkeypatch, synthetic_media, tmp_path, interview_provider):
    source = synthetic_media('one.wav')
    (tmp_path / 'two.wav').write_bytes(source.read_bytes())
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    output = tmp_path / 'out'
    argv = ['--workflow', '--interview', '--input', str(tmp_path), '--output-folder', str(output),
            '--interviewer-name', 'Host', '--interviewee-name', 'Guest', '--stages', 'raw']
    assert main(argv) == 0
    assert len(list((output / 'attributed').rglob('transcription.txt'))) == 2
    assert len(list(output.rglob('transcription.txt'))) == 4


@pytest.mark.parametrize('full,segment,expected', [('Hello world.', 'Helloworld.', False),
    ('A: Hello world.', 'Hello world.', True), ('Hello world.', 'Hello world.', True)])
def test_full_text_comparison_preserves_word_boundaries(full, segment, expected):
    from src.interview_attribution import _complete_turn_text
    assert _complete_turn_text({'text': full, 'segments': [dict(speaker='A', text=segment)]}) == expected


def assert_part_artifacts(output, family, provenance):
    raw = (family / 'transcription.txt').read_bytes().decode('utf-8')
    assert provenance['attribution']['contract'] == 2
    assert provenance['attribution']['part_artifact_reference_base'] == 'output_directory'
    assert state(family)['version'] == 'attributed-interview-v2'
    for part in provenance['parts']:
        path = output / part['raw_transcript']
        assert path.parent == family
        assert path.read_bytes() == raw[part['start']:part['end']].encode('utf-8')
        assert hashlib.sha256(path.read_bytes()).hexdigest() == part['raw_sha256']
        assert state(family)['raw_artifacts'][path.name] == part['raw_sha256']


@pytest.mark.parametrize('change', ['deleted', 'symlink', 'forged_checksum', 'unexpected_path', 'legacy_manifest'])
def test_part_contract_resume_rejects_before_provider(monkeypatch, synthetic_media, tmp_path, interview_provider, change):
    output = tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(synthetic_media(), output, '--stages', 'raw')
    assert main(argv) == 0
    family = next(path.parent for path in (output / 'attributed').rglob('transcription.txt'))
    part = family / 'part-000001_transcription.txt'
    manifest = state(family)
    if change == 'deleted':
        part.unlink()
    elif change == 'symlink':
        part.unlink()
        part.symlink_to(family / 'transcription.txt')
    elif change == 'forged_checksum':
        part.write_bytes(b'forged part')
        manifest['raw_artifacts'][part.name] = hashlib.sha256(part.read_bytes()).hexdigest()
    elif change == 'unexpected_path':
        manifest['raw_artifacts']['../outside.txt'] = manifest['raw_artifacts'].pop(part.name)
    else:
        manifest['version'] = 'attributed-interview-v1'
    (family / 'manifest.json').write_text(json.dumps(manifest))
    before = (family / 'manifest.json').read_bytes()
    calls = interview_provider.audio.transcriptions.create.call_count
    assert main([*argv, '--resume']) == 1
    assert interview_provider.audio.transcriptions.create.call_count == calls
    assert (family / 'manifest.json').read_bytes() == before


def test_v1_family_remains_untouched_while_v2_reuses_cache(monkeypatch, synthetic_media, tmp_path, interview_provider):
    output = tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(synthetic_media(), output, '--stages', 'raw,review')
    assert main(argv) == 0
    current = next(path.parent for path in (output / 'attributed').rglob('transcription.txt'))
    old_fingerprint = _fingerprint({'contract': 1, 'interviewer': 'Synthetic Interviewer',
        'interviewee': 'Synthetic Guest', 'model': DIARIZATION_MODEL, 'mapping': []})
    legacy = current.parent / old_fingerprint
    assert legacy != current
    shutil.copytree(current, legacy)
    old = state(legacy)
    old['version'] = 'attributed-interview-v1'
    payload = json.loads(next(output.rglob('provider_responses.json')).read_text())
    old['binding_sha256'] = _fingerprint({'source': current.parent.name, 'names': old_fingerprint,
                                       'diarization': payload['configuration']})
    provenance = json.loads((legacy / 'provenance.json').read_text())
    provenance['manifest_sha256'] = old_fingerprint
    provenance['attribution']['contract'] = 1
    del provenance['attribution']['part_artifact_reference_base']
    for part in provenance['parts']:
        part['raw_transcript'] = 'attributed/transcription.txt'
        name = f'{part["id"]}_transcription.txt'
        (legacy / name).unlink()
        del old['raw_artifacts'][name]
    (legacy / 'provenance.json').write_text(json.dumps(provenance))
    old['raw_artifacts']['provenance.json'] = hashlib.sha256((legacy / 'provenance.json').read_bytes()).hexdigest()
    # Old derived artifacts are opaque evidence, never accepted for the new contract.
    (legacy / 'manifest.json').write_text(json.dumps(old))
    before = {p.relative_to(legacy): p.read_bytes() for p in legacy.rglob('*') if p.is_file()}
    shutil.rmtree(current)
    audio_calls = interview_provider.audio.transcriptions.create.call_count
    review_calls = interview_provider.chat.completions.create.call_count
    assert main([*argv, '--resume']) == 0
    assert interview_provider.audio.transcriptions.create.call_count == audio_calls
    assert interview_provider.chat.completions.create.call_count > review_calls
    assert before == {p.relative_to(legacy): p.read_bytes() for p in legacy.rglob('*') if p.is_file()}
    assert_part_artifacts(output, current, json.loads((current / 'provenance.json').read_text()))


@pytest.mark.parametrize('ordered', [False, True])
@pytest.mark.parametrize('failure', ['chapter_gate', 'review', 'polish', 'diarization'])
def test_failure_summaries_preserve_completed_attributed_stages(synthetic_media, tmp_path, interview_provider, ordered, failure):
    source, output = synthetic_media(), tmp_path / 'out'
    audio = interview_provider.audio.transcriptions.create.side_effect
    chat = interview_provider.chat.completions.create.side_effect
    if failure == 'chapter_gate':
        interview_provider.high = True
    elif failure == 'diarization':
        def broken_audio(**kw):
            if kw['model'] == DIARIZATION_MODEL:
                raise RuntimeError('PRIVATE_PROVIDER_DETAILS')
            return audio(**kw)
        interview_provider.audio.transcriptions.create.side_effect = broken_audio
    else:
        def broken_chat(**kw):
            body = json.loads(kw['messages'][-1]['content'])
            name = kw['response_format']['json_schema']['name']
            attributed = 'user_confirmed_mapping' in body.get('text', '') or 'unidentified' in body.get('text', '')
            if (failure == 'review' and attributed and name == 'source_grounded_author_review'
                    or failure == 'polish' and any(t['text'] == 'A question?' for t in body.get('turns', []))):
                raise RuntimeError('PRIVATE_PROVIDER_DETAILS')
            return chat(**kw)
        interview_provider.chat.completions.create.side_effect = broken_chat
    options = AuthorOptions(chapter_options=ChapterOptions())
    common = dict(author_options=options, interview_options=InterviewOptions('Host', 'Guest'))
    if ordered:
        manifest = tmp_path / 'parts.json'
        manifest.write_text(json.dumps({'version': 1, 'interview_id': 'synthetic',
            'parts': [{'id': 'part-1', 'path': source.name}]}))
        run = OrderedInterview(manifest, output, enhance=True, options=TranscriptionOptions(),
                               editing_options=EditingOptions(), **common)
        def process():
            return run.process(transcriber=Transcriber(client=interview_provider))
    else:
        run = Pipeline(output, **common)
        def process():
            return run.process(source, transcriber=Transcriber(client=interview_provider), enhance=True)
    with pytest.raises(PipelineError) as caught:
        process()
    stages = caught.value.stages
    assert 'PRIVATE_PROVIDER_DETAILS' not in str(caught.value)
    assert 'attribution' not in stages
    assert stages['attributed_attribution'] == ('failed' if failure == 'diarization' else 'complete')
    if failure != 'diarization':
        family = next(path.parent for path in (output / 'attributed').rglob('transcription.txt'))
        assert state(family)['stages']['transcription']['status'] == 'complete'
    if failure in ('chapter_gate', 'review'):
        assert stages['attributed_enhancement'] == 'complete'
        assert stages['attributed_author_review'] == ('complete' if failure == 'chapter_gate' else 'failed')
    if failure == 'chapter_gate':
        assert stages['attributed_chapters'] == 'failed'
        assert state(family)['stages']['author_review']['status'] == 'complete'
        run.resume = True
        with pytest.raises(PipelineError) as resumed:
            process()
        assert resumed.value.stages['attributed_attribution'] == 'skipped'
        assert resumed.value.stages['attributed_enhancement'] == 'skipped'
        assert resumed.value.stages['attributed_author_review'] == 'skipped'
        assert resumed.value.stages['attributed_chapters'] == 'failed'
    if failure == 'polish':
        assert stages['attributed_enhancement'] == 'failed'
    assert not list((output / 'attributed').rglob('chapter_drafts.json'))


@pytest.mark.parametrize('mapping', [(), ('1:1:A=interviewer', '1:1:B=interviewee')])
def test_attributed_chapter_banner_includes_mapping_uncertainty(monkeypatch, synthetic_media, tmp_path, interview_provider, mapping):
    output = tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    extras = [item for value in mapping for item in ('--speaker-map', value)]
    assert main(args(synthetic_media(), output, '--stages', 'raw,review', '--chapters', 'interview', *extras)) == 0
    family = next(path.parent for path in (output / 'attributed').rglob('transcription.txt'))
    text = artifact(family, 'chapters', 'chapter_interview.txt').read_text()
    assert 'Names and roles appear only for user-confirmed mappings; unmapped speakers remain unidentified.' in text
    assert 'Speaker roles are unassigned.' not in text
    assert 'Speaker 1:1:C: Speaker C; identity evidence: unidentified' in text
    assert ('identity evidence: user_confirmed_mapping' in text) == bool(mapping)


def test_part_files_preserve_unicode_and_crlf_exact_bytes(monkeypatch, synthetic_media, tmp_path, interview_provider):
    audio = interview_provider.audio.transcriptions.create.side_effect
    def unicode_audio(**kw):
        response = audio(**kw)
        if isinstance(response, dict):
            response['segments'][0]['text'] = 'A café question?\r\n第二行'
            response['text'] = '\n'.join(f'{s["speaker"]}: {s["text"]}' for s in response['segments'])
        return response
    interview_provider.audio.transcriptions.create.side_effect = unicode_audio
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    output = tmp_path / 'out'
    argv = args(synthetic_media(), output, '--stages', 'raw')
    assert main(argv) == 0
    family = next(path.parent for path in (output / 'attributed').rglob('transcription.txt'))
    assert_part_artifacts(output, family, json.loads((family / 'provenance.json').read_text()))
    assert 'A café question?\r\n第二行'.encode() in (family / 'part-000001_transcription.txt').read_bytes()
    calls = interview_provider.audio.transcriptions.create.call_count
    assert main([*argv, '--resume']) == 0
    assert interview_provider.audio.transcriptions.create.call_count == calls


@pytest.mark.parametrize('ordered', [False, True])
def test_cli_failed_chapter_reports_completed_attribution(monkeypatch, synthetic_media, tmp_path, interview_provider, capsys, ordered):
    interview_provider.high = True
    source, output = synthetic_media(), tmp_path / 'out'
    monkeypatch.setattr('src.cli.Transcriber', lambda **kw: Transcriber(client=interview_provider, **kw))
    argv = args(source, output, '--stages', 'raw,review', '--chapters', 'interview')
    if ordered:
        manifest = tmp_path / 'parts.json'
        manifest.write_text(json.dumps({'version': 1, 'interview_id': 'synthetic',
            'parts': [{'id': 'part-1', 'path': source.name}]}))
        argv[argv.index('--input'):argv.index('--input') + 2] = ['--interview-manifest', str(manifest)]
    assert entrypoint(argv) == 1
    logs = capsys.readouterr()
    events = [json.loads(line) for line in logs.out.splitlines() if line.startswith('{')]
    failure = next(event for event in events if event['status'] == 'failed' and 'stages' in event)
    assert failure['stages']['attributed_attribution'] == 'complete'
    assert failure['stages']['attributed_author_review'] == 'complete'
    assert failure['stages']['attributed_chapters'] == 'failed'
    assert 'attribution' not in failure['stages']
    for private in ('Synthetic Interviewer', 'Synthetic Guest', 'A question?', str(tmp_path), source.name):
        assert private not in logs.out + logs.err
    assert not list(output.rglob('chapter_drafts.json'))
