"""Synthetic-only ordered interview, no provider network or private recordings."""

import hashlib
import json
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from src.author_review import ReviewOptions
from src.author_workflow import AuthorOptions
from src.chapters import ChapterOptions
from src.cli import main
from src.model_config import EditingOptions, TranscriptionOptions
from src.ordered_interview import OrderedInterview, load_interview
from src.pipeline import PipelineError
from src.source_provenance import references, validate_provenance
from src.transcriber import Transcriber


@pytest.fixture
def interview_provider(provider):
    def respond(**kwargs):
        supplied = json.loads(kwargs["messages"][-1]["content"])
        name = kwargs["response_format"]["json_schema"]["name"]
        if name == "faithful_transcript_edit":
            body = {
                "chunk_index": supplied["chunk_index"],
                "text": supplied["text"],
                "speaker_uncertain": True,
            }
        elif name == "source_grounded_author_review":
            supplied['text'] = ''.join(p['text'] for p in supplied['evidence_pieces'])
            findings = []
            if provider.priority:
                code = "serious_allegation" if provider.priority == "high" else "personal_opinion"
                findings = [
                    {
                        "category": "allegation" if provider.priority == "high" else "criticism",
                        "severity": provider.priority,
                        "reason_code": code,
                        "excerpt": supplied["text"],
                        "piece_ids": [p["piece_id"] for p in supplied["evidence_pieces"]],
                    }
                ]
            body = {
                "chunk_index": supplied["chunk_index"],
                "fully_reviewed": True,
                "contract_version": supplied["contract_version"],
                "reviewed_piece_ids": supplied["core_piece_ids"],
                "findings": findings,
            }
        else:
            body = {
                "chunk_index": supplied["chunk_index"],
                "passages": [
                    {
                        "unit_ids": [u["unit_id"]],
                        "text": u["text"].strip(),
                        "kind": "verbatim_excerpt",
                    }
                    for u in supplied["source_units"]
                ],
                "coverage_omissions": [],
            }
        return SimpleNamespace(
            choices=[
                SimpleNamespace(
                    finish_reason="stop",
                    message=SimpleNamespace(
                        content=json.dumps(body, ensure_ascii=False), refusal=None
                    ),
                )
            ]
        )

    provider.priority = None
    provider.chat.completions.create.side_effect = respond
    return provider


def manifest(tmp_path, sources):
    path = tmp_path / "synthetic-interview.json"
    document = {
        "version": 1,
        "interview_id": "synthetic-interview",
        "parts": [
            {"id": f"part-{i}", "path": str(p), "media_type": "auto"}
            for i, p in enumerate(sources, 1)
        ],
    }
    path.write_text(json.dumps(document))
    return path


def workflow(path, output, *, resume=False, review=False, styles=(), enhance=False, **kwargs):
    return OrderedInterview(
        path,
        output,
        options=kwargs.get("options", TranscriptionOptions()),
        editing_options=kwargs.get("editing_options", EditingOptions()),
        resume=resume,
        enhance=enhance,
        author_options=AuthorOptions(
            review=review or bool(styles),
            review_options=kwargs.get("review_options", ReviewOptions()),
            chapter_options=ChapterOptions(styles=styles, chunk_bytes=64) if styles else None,
            allow_unresolved_high=kwargs.get("override", False),
        ),
    )


def read_artifact(run, stage, filename):
    state = json.loads((run.job / "manifest.json").read_text())
    target = next(
        run.job / p for p in state["stages"][stage]["artifacts"] if p.endswith("/" + filename)
    )
    return target


def test_explicit_order_mixed_formats_exact_raw_and_provenance(
    synthetic_media, tmp_path, interview_provider
):
    sources = [
        synthetic_media("Part10.mp4"),
        synthetic_media("Part2.wav"),
        synthetic_media("continuation.m4a"),
    ]
    source_hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in sources]
    texts = ["First न 👩🏽‍💻 +−²\r\n", "Second € = 3.\n", "Third testimony."]
    interview_provider.audio.transcriptions.create.side_effect = [
        SimpleNamespace(text=t) for t in texts
    ]
    run = workflow(
        manifest(tmp_path, sources),
        tmp_path / "output",
        review=True,
        styles=("interview", "narrative"),
        enhance=True,
        review_options=ReviewOptions(chunk_bytes=64),
    )
    identity, summary = run.process(transcriber=Transcriber(client=interview_provider))
    assert identity == run.binding and summary["combined_raw"] == "complete"
    assert summary["author_review"] == summary["chapters"] == "complete"
    raw = "\n\n".join(texts)
    assert (run.job / "transcription.txt").read_bytes() == raw.encode()
    assert [hashlib.sha256(p.read_bytes()).hexdigest() for p in sources] == source_hashes
    provenance = json.loads((run.job / "provenance.json").read_bytes())
    validate_provenance(raw, provenance)
    assert [p["path"] for p in provenance["parts"]] == [str(p.resolve()) for p in sources]
    for part, text in zip(provenance["parts"], texts):
        assert raw[part["start"] : part["end"]] == text
        assert (run.output / part["raw_transcript"]).read_bytes() == text.encode()
        assert references(provenance, part["start"], part["end"])[0]["local_start"] == 0
    assert references(provenance, 0, len(raw))[1]["kind"] == "separator"
    report = json.loads(read_artifact(run, "author_review", "review_report.json").read_bytes())
    assert report["raw_sha256"] == provenance["raw_sha256"]
    assert report["recording_provenance"] == provenance
    assert report["coverage"]["reviewed_characters"] == len(raw)
    wb = load_workbook(read_artifact(run, "author_review", "review_report.xlsx"))
    assert wb["Recording parts"].cell(2, 3).value == str(sources[0].resolve())
    assert wb["Recording parts"].max_row == 4
    assert wb["Recording references"].max_row > len(report["segments"])
    chapters = json.loads(read_artifact(run, "chapters", "chapter_drafts.json").read_bytes())
    assert set(chapters["chapters"]) == {
        "interview",
        "narrative",
    }  # one set for interview, not per file
    for chapter in chapters["chapters"].values():
        assert chapter["coverage"]["accounted_characters"] == len(raw)
        for passage in chapter["passages"]:
            assert passage["recording_refs"] == references(
                provenance, passage["start"], passage["end"]
            )
            assert passage["text"] == raw[passage["quote_start"] : passage["quote_end"]]
    assert "Recording part 1" in read_artifact(run, "chapters", "chapter_interview.txt").read_text()
    # Private source paths are NEVER sent to editing/review/chapter models.
    calls = str(interview_provider.chat.completions.create.call_args_list)
    assert str(tmp_path) not in calls
    assert "Part10.mp4" not in calls
    assert run.job.stat().st_mode & 0o777 == 0o700
    assert all(p.stat().st_mode & 0o777 == 0o600 for p in run.job.iterdir() if p.is_file())


@pytest.mark.parametrize(
    "mutation",
    [
        "version",
        "bool_version",
        "id",
        "empty_parts",
        "duplicate_id",
        "same_file",
        "hardlink",
        "missing",
        "empty",
        "unsupported",
        "media_type",
        "url",
        "empty_path",
        "symlink",
        "parent_symlink",
        "unknown",
        "explicit_order",
    ],
)
def test_manifest_validation_before_provider(mutation, tmp_path, monkeypatch, capsys):
    source = tmp_path / "SYNTHETIC_PRIVATE.wav"
    source.write_bytes(b"nonempty synthetic")
    other = tmp_path / "synthetic-second.wav"
    other.write_bytes(b"other nonempty synthetic")
    path = manifest(tmp_path, [source, other])
    doc = json.loads(path.read_text())
    if mutation == "version":
        doc["version"] = 0
    elif mutation == "bool_version":
        doc["version"] = True
    elif mutation == "id":
        doc["interview_id"] = ""
    elif mutation == "empty_parts":
        doc["parts"] = []
    elif mutation == "duplicate_id":
        doc["parts"][1]["id"] = doc["parts"][0]["id"]
    elif mutation == "same_file":
        doc["parts"][1]["path"] = str(tmp_path / "." / source.name)
    elif mutation == "hardlink":
        other.unlink()
        other.hardlink_to(source)
    elif mutation == "missing":
        other.unlink()
    elif mutation == "empty":
        other.write_bytes(b"")
    elif mutation == "unsupported":
        doc["parts"][1]["path"] = "synthetic.txt"
    elif mutation == "media_type":
        doc["parts"][0]["media_type"] = "video"
    elif mutation == "url":
        doc["parts"][0]["path"] = "https://example.invalid/SYNTHETIC_PRIVATE.wav"
    elif mutation == "empty_path":
        doc["parts"][0]["path"] = ""
    elif mutation == "symlink":
        other.unlink()
        other.symlink_to(source)
    elif mutation == "parent_symlink":
        alias = tmp_path / "alias"
        alias.symlink_to(tmp_path, target_is_directory=True)
        doc["parts"][1]["path"] = str(alias / other.name)
    elif mutation == "unknown":
        doc["SYNTHETIC_PRIVATE"] = True
    else:
        doc["parts"][0]["order"] = 2
    path.write_text(json.dumps(doc))
    factory = MagicMock(side_effect=AssertionError("No provider initialization"))
    monkeypatch.setattr("src.cli.Transcriber", factory)
    assert (
        main(
            [
                "--workflow",
                "--interview-manifest",
                str(path),
                "--stages",
                "raw",
                "--output-folder",
                str(tmp_path / "output"),
            ]
        )
        == 1
    )
    factory.assert_not_called()
    logs = capsys.readouterr()
    assert "SYNTHETIC_PRIVATE" not in logs.out + logs.err and str(tmp_path) not in logs.out


@pytest.mark.parametrize("content", ['{"version":1,"version":1}', "{invalid", "[]", "null"])
def test_duplicate_json_keys_and_invalid_json(tmp_path, content):
    path = tmp_path / "synthetic.json"
    path.write_text(content)
    with pytest.raises(PipelineError, match="Invalid interview manifest"):
        load_interview(path)


@pytest.mark.parametrize("failure", ["corrupt", "no_audio"])
def test_all_local_conversion_before_any_paid_work(synthetic_media, tmp_path, provider, failure):
    first = synthetic_media("synthetic-first.wav")
    if failure == "no_audio":
        second = synthetic_media("synthetic-second.mp4", audio=False)
    else:
        second = tmp_path / "synthetic-second.wav"
        second.write_bytes(b"corrupt synthetic media")
    run = workflow(manifest(tmp_path, [first, second]), tmp_path / "output")
    with pytest.raises(PipelineError):
        run.process(transcriber=Transcriber(client=provider))
    provider.audio.transcriptions.create.assert_not_called()
    assert not run.job.exists()
    assert any(run.output.rglob("audio.wav"))


@pytest.mark.parametrize("failure", ["error", "empty"])
def test_part_failure_no_combined_or_author_work_safe_resume(
    synthetic_media, tmp_path, interview_provider, failure
):
    sources = [synthetic_media("synthetic-a.wav"), synthetic_media("synthetic-b.mp4")]
    path = manifest(tmp_path, sources)
    run = workflow(path, tmp_path / "output", review=True, enhance=True, styles=("interview",))
    bad = (
        RuntimeError("SYNTHETIC_PRIVATE_PROVIDER_ERROR")
        if failure == "error"
        else SimpleNamespace(text=" \r\n")
    )
    interview_provider.audio.transcriptions.create.side_effect = [
        SimpleNamespace(text="First."),
        bad,
    ]
    with pytest.raises(PipelineError) as error:
        run.process(transcriber=Transcriber(client=interview_provider))
    assert "SYNTHETIC_PRIVATE" not in str(error.value)
    assert not run.job.exists()
    assert len(list(run.output.rglob("transcription.txt"))) == 1
    interview_provider.chat.completions.create.assert_not_called()
    interview_provider.audio.transcriptions.create.side_effect = [SimpleNamespace(text="Second.")]
    resumed = workflow(
        path, run.output, resume=True, review=True, enhance=True, styles=("interview",)
    )
    resumed.process(transcriber=Transcriber(client=interview_provider))
    assert interview_provider.audio.transcriptions.create.call_count == 3
    assert (resumed.job / "transcription.txt").read_text() == "First.\n\nSecond."


@pytest.mark.parametrize("change", ["reorder", "change_source", "add", "remove", "rename_id"])
def test_manifest_changes_new_combined_reuse_verified_parts(
    synthetic_media, tmp_path, interview_provider, change
):
    sources = [synthetic_media("Part2.wav"), synthetic_media("Part10.mp4")]
    path = manifest(tmp_path, sources)
    output = tmp_path / "output"
    interview_provider.audio.transcriptions.create.side_effect = [
        SimpleNamespace(text="First."),
        SimpleNamespace(text="Second."),
    ]
    run = workflow(path, output, review=True)
    run.process(transcriber=Transcriber(client=interview_provider))
    old_raw = (run.job / "transcription.txt").read_bytes()
    old_report = read_artifact(run, "author_review", "review_report.json").read_bytes()
    doc = json.loads(path.read_text())
    extra_calls = 0
    expected = "First.\n\nSecond."
    if change == "reorder":
        doc["parts"].reverse()
        expected = "Second.\n\nFirst."
    elif change == "change_source":
        sources[1].write_bytes(sources[1].read_bytes() + b"synthetic harmless trailer")
        extra_calls = 1
        expected = "First.\n\nNew."
    elif change == "add":
        new = synthetic_media("unnumbered.wav")
        doc["parts"].insert(1, {"id": "extra", "path": str(new)})
        extra_calls = 1
        expected = "First.\n\nNew.\n\nSecond."
    elif change == "remove":
        doc["parts"].pop(0)
        expected = "Second."
    else:
        doc["parts"][0]["id"] = "new-id"
    path.write_text(json.dumps(doc))
    interview_provider.audio.transcriptions.create.side_effect = [SimpleNamespace(text="New.")]
    calls_before = interview_provider.chat.completions.create.call_count
    new_run = workflow(path, output, resume=True, review=True)
    new_run.process(transcriber=Transcriber(client=interview_provider))
    assert run.job != new_run.job
    assert (new_run.job / "transcription.txt").read_text() == expected
    assert interview_provider.audio.transcriptions.create.call_count == 2 + extra_calls
    assert interview_provider.chat.completions.create.call_count > calls_before
    assert (run.job / "transcription.txt").read_bytes() == old_raw
    assert read_artifact(run, "author_review", "review_report.json").read_bytes() == old_report


def test_verified_resume_skips_every_stage(synthetic_media, tmp_path, interview_provider):
    path = manifest(tmp_path, [synthetic_media("synthetic.wav")])
    first = workflow(
        path, tmp_path / "output", review=True, styles=("interview", "narrative"), enhance=True
    )
    transcriber = Transcriber(client=interview_provider)
    first.process(transcriber=transcriber)
    calls = interview_provider.mock_calls[:]
    resumed = workflow(
        path,
        first.output,
        resume=True,
        review=True,
        styles=("interview", "narrative"),
        enhance=True,
    )
    _, summary = resumed.process(transcriber=transcriber)
    assert summary == {
        "parts": "complete",
        "combined_raw": "skipped",
        "enhancement": "skipped",
        "author_review": "skipped",
        "chapters": "skipped",
    }
    assert interview_provider.mock_calls == calls


@pytest.mark.parametrize(
    "target",
    [
        "part_raw",
        "part_audio",
        "part_manifest",
        "provenance",
        "combined",
        "combined_manifest",
        "boundaries",
        "review_json",
        "review_xlsx",
        "chapter_json",
        "polish",
    ],
)
def test_cache_collisions_detected_before_asr(
    synthetic_media, tmp_path, interview_provider, target
):
    path = manifest(tmp_path, [synthetic_media("synthetic.wav")])
    first = workflow(path, tmp_path / "output", review=True, styles=("interview",), enhance=True)
    first.process(transcriber=Transcriber(client=interview_provider))
    part, directory, identity = first.parts[0]
    mapping = {
        "part_raw": directory / identity / "transcription.txt",
        "part_audio": directory / identity / "audio.wav",
        "part_manifest": directory / identity / "manifest.json",
        "provenance": first.job / "provenance.json",
        "combined": first.job / "transcription.txt",
        "combined_manifest": first.job / "manifest.json",
        "boundaries": first.job / "part_boundaries.txt",
        "review_json": read_artifact(first, "author_review", "review_report.json"),
        "review_xlsx": read_artifact(first, "author_review", "review_report.xlsx"),
        "chapter_json": read_artifact(first, "chapters", "chapter_drafts.json"),
        "polish": first.job / "derivative_readability.txt",
    }
    mapping[target].write_bytes(b"SYNTHETIC_PRIVATE_TAMPERED")
    calls = interview_provider.mock_calls[:]
    # Add a fresh ASR part so the test proves all existing parts are preflighted first.
    if target.startswith("part_"):
        doc = json.loads(path.read_text())
        doc["parts"].insert(0, {"id": "new", "path": str(synthetic_media("synthetic-new.wav"))})
        path.write_text(json.dumps(doc))
    new = workflow(
        path, first.output, resume=True, review=True, styles=("interview",), enhance=True
    )
    with pytest.raises(PipelineError):
        new.process(transcriber=Transcriber(client=interview_provider))
    assert interview_provider.mock_calls == calls
    assert mapping[target].read_bytes() == b"SYNTHETIC_PRIVATE_TAMPERED"


@pytest.mark.parametrize("change", ["asr", "editing", "review"])
def test_settings_change_new_generation(synthetic_media, tmp_path, interview_provider, change):
    path = manifest(tmp_path, [synthetic_media("synthetic.wav")])
    first = workflow(path, tmp_path / "output", review=True, enhance=True)
    first.process(transcriber=Transcriber(client=interview_provider))
    kwargs = {}
    asr, editing = TranscriptionOptions(), EditingOptions()
    if change == "asr":
        asr = TranscriptionOptions(context="Synthetic hint.")
    elif change == "editing":
        editing = EditingOptions(model="gpt-6.1-sol")
    else:
        kwargs["review_options"] = ReviewOptions(model="gpt-6.1-sol")
    new = workflow(
        path,
        first.output,
        resume=True,
        review=True,
        enhance=True,
        options=asr,
        editing_options=editing,
        **kwargs,
    )
    new.process(
        transcriber=Transcriber(client=interview_provider, options=asr, editing_options=editing)
    )
    assert first.job != new.job
    assert interview_provider.audio.transcriptions.create.call_count == (
        2 if change == "asr" else 1
    )


def test_high_gate_combined_boundary_finding(synthetic_media, tmp_path, interview_provider):
    path = manifest(
        tmp_path, [synthetic_media("synthetic-a.wav"), synthetic_media("synthetic-b.wav")]
    )
    interview_provider.audio.transcriptions.create.side_effect = [
        SimpleNamespace(text="First."),
        SimpleNamespace(text="Second."),
    ]
    interview_provider.priority = "high"
    run = workflow(path, tmp_path / "output", review=True, styles=("interview",))
    with pytest.raises(PipelineError, match="unresolved high"):
        run.process(transcriber=Transcriber(client=interview_provider))
    report = json.loads(read_artifact(run, "author_review", "review_report.json").read_bytes())
    assert len(report["findings"]) == 1
    assert [r["kind"] for r in report["findings"][0]["recording_refs"]] == [
        "recording",
        "separator",
        "recording",
    ]
    assert not list(run.job.rglob("chapter_drafts.json"))
    override = workflow(
        path, run.output, resume=True, review=True, styles=("interview",), override=True
    )
    override.process(transcriber=Transcriber(client=interview_provider))
    draft = json.loads(read_artifact(override, "chapters", "chapter_drafts.json").read_bytes())
    assert draft["unresolved_high_finding_ids"] and draft["human_review_required"]
    assert any("DRAFT OVERRIDE" in w for w in draft["chapters"]["interview"]["warnings"])


@pytest.mark.parametrize(
    "args",
    [
        ["--interview-manifest", "unused"],
        ["--workflow", "--interview-manifest", "", "--stages", "raw"],
        ["--workflow", "--interview-manifest", "unused", "--input", "unused"],
        ["--workflow", "--interview-manifest", "unused", "--media-type", "video"],
        ["--workflow", "--interview-manifest", "unused", "--extract-only"],
    ],
)
def test_cli_usage_validation(monkeypatch, args):
    factory = MagicMock(side_effect=AssertionError())
    monkeypatch.setattr("src.cli.Transcriber", factory)
    try:
        assert main(args) == 1
    except SystemExit as error:
        assert error.code == 2
    factory.assert_not_called()


def test_cli_ordered_one_interview_private_logs(
    synthetic_media, tmp_path, interview_provider, monkeypatch, capsys
):
    path = manifest(
        tmp_path, [synthetic_media("SYNTHETIC_PRIVATE.mp4"), synthetic_media("Part2.wav")]
    )
    interview_provider.audio.transcriptions.create.side_effect = [
        SimpleNamespace(text="SYNTHETIC_PRIVATE first."),
        SimpleNamespace(text="Second."),
    ]
    monkeypatch.setattr(
        "src.cli.Transcriber", lambda **kwargs: Transcriber(client=interview_provider, **kwargs)
    )
    assert (
        main(
            [
                "--workflow",
                "--interview-manifest",
                str(path),
                "--output-folder",
                str(tmp_path / "output"),
                "--stages",
                "raw,polish,review",
                "--chapters",
                "both",
            ]
        )
        == 0
    )
    logs = capsys.readouterr().out
    assert "SYNTHETIC_PRIVATE" not in logs and str(tmp_path) not in logs
    assert json.loads(logs.splitlines()[-1]) == {
        "status": "summary", "processed": 1, "completed": 1, "failed": 0,
        "selected": 1, "finished": 1,
    }


def test_installed_entrypoint_help_outside_checkout(tmp_path):
    # Regression for pilot launcher import failure: exercise the installed executable.
    launcher = Path(sys.executable).parent / "voice-transcribe"
    result = subprocess.run(
        [str(launcher), "--help"], cwd=tmp_path, capture_output=True, timeout=30
    )
    assert result.returncode == 0
    assert b"--interview-manifest" in result.stdout


def test_relative_paths_semantic_manifest_identity(synthetic_media, tmp_path, provider):
    source = synthetic_media("synthetic.wav")
    path = manifest(tmp_path, [source])
    first = workflow(path, tmp_path / "output")
    first.process(transcriber=Transcriber(client=provider))
    document = json.loads(path.read_text())
    document["parts"][0]["path"] = "./synthetic.wav"
    path.write_text(json.dumps(document, indent=4, sort_keys=True))
    second = workflow(path, first.output, resume=True)
    assert first.binding == second.binding
    second.process(transcriber=Transcriber(client=provider))
    assert provider.audio.transcriptions.create.call_count == 1


@pytest.mark.parametrize("stage", ["author_review", "chapters"])
def test_invalid_bound_refs_even_with_updated_artifact_hash(
    synthetic_media, tmp_path, interview_provider, stage
):
    from src.private_output import digest

    path = manifest(tmp_path, [synthetic_media("synthetic.wav")])
    run = workflow(path, tmp_path / "output", review=True, styles=("interview",))
    run.process(transcriber=Transcriber(client=interview_provider))
    name = "review_report.json" if stage == "author_review" else "chapter_drafts.json"
    target = read_artifact(run, stage, name)
    document = json.loads(target.read_text())
    if stage == "author_review":
        document["recording_segment_refs"] = {}
    else:
        document["chapters"]["interview"]["passages"][0]["recording_refs"] = []
    target.write_text(json.dumps(document))
    state = json.loads((run.job / "manifest.json").read_text())
    state["stages"][stage]["artifacts"][str(target.relative_to(run.job))] = digest(target)
    (run.job / "manifest.json").write_text(json.dumps(state))
    calls = interview_provider.mock_calls[:]
    new = workflow(path, run.output, resume=True, review=True, styles=("interview",))
    with pytest.raises(PipelineError):
        new.process(transcriber=Transcriber(client=interview_provider))
    assert calls == interview_provider.mock_calls


@pytest.mark.parametrize(
    "kind", ["output_parent", "lock", "manifest_version", "stray_raw", "symlink_raw"]
)
def test_preflight_output_conflicts(synthetic_media, tmp_path, provider, kind):
    path = manifest(tmp_path, [synthetic_media("synthetic.wav")])
    output = tmp_path / "output"
    if kind == "output_parent":
        output.mkdir()
        (output / "interviews").write_bytes(b"synthetic collision")
        with pytest.raises((OSError, ValueError)):
            workflow(path, output)
    else:
        run = workflow(path, output)
        if kind == "lock":
            (output / ".interview.lock").touch()
        elif kind == "stray_raw":
            run.job.mkdir()
            (run.job / "transcription.txt").write_text("Synthetic conflicting text.")
        else:
            run.process(transcriber=Transcriber(client=provider))
            provider.reset_mock()
            if kind == "manifest_version":
                part, directory, identity = run.parts[0]
                state_path = directory / identity / "manifest.json"
                state = json.loads(state_path.read_text())
                state["version"] = 1
                state_path.write_text(json.dumps(state))
            else:
                raw = run.job / "transcription.txt"
                raw.unlink()
                raw.symlink_to(path)
        with pytest.raises(PipelineError):
            workflow(path, output, resume=True).process(transcriber=Transcriber(client=provider))
    provider.audio.transcriptions.create.assert_not_called()


def test_source_change_during_asr_blocks_combined(synthetic_media, tmp_path, provider):
    source = synthetic_media("synthetic.wav")
    run = workflow(manifest(tmp_path, [source]), tmp_path / "output", review=True)

    def transcribe(**kwargs):
        source.write_bytes(source.read_bytes() + b"synthetic changed trailer")
        return SimpleNamespace(text="Synthetic transcript.")

    provider.audio.transcriptions.create.side_effect = transcribe
    with pytest.raises(PipelineError):
        run.process(transcriber=Transcriber(client=provider))
    assert not run.job.exists()
    provider.chat.completions.create.assert_not_called()


def test_polish_fidelity_failure_retains_combined_and_raw_parts(
    synthetic_media, tmp_path, provider
):
    run = workflow(
        manifest(tmp_path, [synthetic_media("synthetic.wav")]),
        tmp_path / "output",
        review=True,
        enhance=True,
    )
    provider.chat.completions.create.return_value.choices[0].message.content = json.dumps(
        {"chunk_index": 1, "text": "Invented words.", "speaker_uncertain": True}
    )
    with pytest.raises(PipelineError):
        run.process(transcriber=Transcriber(client=provider))
    assert (run.job / "transcription.txt").read_text() == "Synthetic transcript."
    assert not (run.job / "derivative_readability.txt").exists()
    assert not list(run.job.rglob("review_report.json"))


def test_report_export_rejects_wrong_local_offsets(synthetic_media, tmp_path, interview_provider):
    from src.review_export import export_review, ReviewExportError

    run = workflow(
        manifest(tmp_path, [synthetic_media("synthetic.wav")]), tmp_path / "output", review=True
    )
    run.process(transcriber=Transcriber(client=interview_provider))
    report = json.loads(read_artifact(run, "author_review", "review_report.json").read_text())
    segment = next(iter(report["recording_segment_refs"].values()))
    segment[0]["local_start"] = 123
    with pytest.raises(ReviewExportError):
        export_review(report, tmp_path / "synthetic-invalid.xlsx")
    assert not (tmp_path / "synthetic-invalid.xlsx").exists()


def test_polish_prompt_change_reuses_asr_creates_new_generation(
    synthetic_media, tmp_path, interview_provider, monkeypatch
):
    from src import transcriber as asr_engine

    path = manifest(tmp_path, [synthetic_media("synthetic.wav")])
    first = workflow(path, tmp_path / "output", enhance=True, review=True)
    first.process(transcriber=Transcriber(client=interview_provider))
    old = (first.job / "derivative_readability.txt").read_bytes()
    monkeypatch.setattr(
        asr_engine, "FAITHFUL_INSTRUCTION", asr_engine.FAITHFUL_INSTRUCTION + " Synthetic revision."
    )
    new = workflow(path, first.output, resume=True, enhance=True, review=True)
    new.process(transcriber=Transcriber(client=interview_provider))
    assert new.binding != first.binding
    assert interview_provider.audio.transcriptions.create.call_count == 1
    assert (first.job / "derivative_readability.txt").read_bytes() == old


def test_later_part_changed_during_first_asr_is_not_transcribed(
    synthetic_media, tmp_path, provider
):
    sources = [synthetic_media("synthetic-first.wav"), synthetic_media("synthetic-later.wav")]
    run = workflow(manifest(tmp_path, sources), tmp_path / "output", review=True)

    def change_later(**kwargs):
        sources[1].write_bytes(sources[1].read_bytes() + b"synthetic changed trailer")
        return SimpleNamespace(text="First completed synthetic part.")

    provider.audio.transcriptions.create.side_effect = change_later
    with pytest.raises(PipelineError):
        run.process(transcriber=Transcriber(client=provider))
    assert provider.audio.transcriptions.create.call_count == 1
    assert not run.job.exists()
    provider.chat.completions.create.assert_not_called()
    part, directory, identity = run.parts[0]
    assert (
        directory / identity / "transcription.txt"
    ).read_text() == "First completed synthetic part."
    # The second part has only its initially validated conversion job, not a job
    # created for changed, unapproved source bytes.
    part, directory, identity = run.parts[1]
    assert list(directory.iterdir()) == [directory / identity]
    assert not (directory / identity / "transcription.txt").exists()


@pytest.mark.parametrize("alteration", ["missing", "incorrect"])
def test_cached_chapter_findings_references_are_revalidated(
    synthetic_media, tmp_path, interview_provider, alteration
):
    from src.private_output import digest

    interview_provider.priority = "low"
    path = manifest(tmp_path, [synthetic_media("synthetic.wav")])
    run = workflow(path, tmp_path / "output", review=True, styles=("interview",))
    run.process(transcriber=Transcriber(client=interview_provider))
    target = read_artifact(run, "chapters", "chapter_drafts.json")
    document = json.loads(target.read_bytes())
    assert document["findings"]
    if alteration == "missing":
        document["findings"][0].pop("recording_refs")
    else:
        document["findings"][0]["recording_refs"][0]["part_id"] = "synthetic-wrong-part"
    target.write_text(json.dumps(document))
    state = json.loads((run.job / "manifest.json").read_bytes())
    state["stages"]["chapters"]["artifacts"][str(target.relative_to(run.job))] = digest(target)
    (run.job / "manifest.json").write_text(json.dumps(state))
    calls = interview_provider.mock_calls[:]
    with pytest.raises(PipelineError):
        workflow(path, run.output, resume=True, review=True, styles=("interview",)).process(
            transcriber=Transcriber(client=interview_provider)
        )
    assert interview_provider.mock_calls == calls


@pytest.mark.parametrize(
    "failure,expected",
    [
        ("fidelity", "changed, invented, omitted, or reordered words or symbols"),
        ("provider", "API/model access, quota, and network connectivity"),
    ],
)
def test_polish_failure_preserves_safe_retry_diagnostic(
    synthetic_media, tmp_path, provider, failure, expected
):
    run = workflow(
        manifest(tmp_path, [synthetic_media("synthetic.wav")]),
        tmp_path / "output",
        enhance=True,
        review=True,
    )
    if failure == "fidelity":
        provider.chat.completions.create.return_value.choices[0].message.content = json.dumps(
            {"chunk_index": 1, "text": "Invented words.", "speaker_uncertain": True}
        )
    else:
        provider.chat.completions.create.side_effect = RuntimeError(
            "SYNTHETIC_PRIVATE_PROVIDER_ERROR"
        )
    with pytest.raises(PipelineError, match=expected) as error:
        run.process(transcriber=Transcriber(client=provider))
    assert "SYNTHETIC_PRIVATE" not in str(error.value)
    assert "Check local media" not in str(error.value)
    assert (run.job / "transcription.txt").read_text() == "Synthetic transcript."
    assert not (run.job / "derivative_readability.txt").exists()


def test_windows_part_combination_and_resume_preserve_exact_bytes(
    synthetic_media, tmp_path, interview_provider, monkeypatch
):
    import os

    real_fdopen = os.fdopen

    def windows_fdopen(fd, mode, **kwargs):
        if "b" not in mode and kwargs.get("newline") is None:
            kwargs["newline"] = "\r\n"
        return real_fdopen(fd, mode, **kwargs)

    monkeypatch.setattr("src.private_output.os.fdopen", windows_fdopen)
    texts = ["First न +²\r\nLF\nCR\r", "Second 👩🏽‍💻\n\r\n"]
    interview_provider.audio.transcriptions.create.side_effect = [
        SimpleNamespace(text=t) for t in texts
    ]
    path = manifest(
        tmp_path, [synthetic_media("synthetic-a.wav"), synthetic_media("synthetic-b.mp4")]
    )
    run = workflow(path, tmp_path / "output", review=True, enhance=True, styles=("interview",))
    run.process(transcriber=Transcriber(client=interview_provider))
    combined = "\n\n".join(texts).encode()
    assert (run.job / "transcription.txt").read_bytes() == combined
    provenance = json.loads((run.job / "provenance.json").read_bytes())
    assert provenance["raw_sha256"] == hashlib.sha256(combined).hexdigest()
    for part, text in zip(provenance["parts"], texts):
        assert (run.output / part["raw_transcript"]).read_bytes() == text.encode()
    calls = interview_provider.mock_calls[:]
    workflow(
        path, run.output, resume=True, review=True, enhance=True, styles=("interview",)
    ).process(transcriber=Transcriber(client=interview_provider))
    assert interview_provider.mock_calls == calls


@pytest.mark.parametrize("failure", ["fidelity", "provider", "unexpected", "subclass"])
def test_cli_polish_failure_allowlist_preserves_guidance_and_privacy(
    synthetic_media, tmp_path, provider, monkeypatch, capsys, failure
):
    from src.transcriber import TranscriptionError

    class UntrustedTranscriptionError(TranscriptionError):
        pass

    transcriber = Transcriber(client=provider)
    if failure == "fidelity":
        provider.chat.completions.create.return_value.choices[0].message.content = json.dumps(
            {"chunk_index": 1, "text": "Invented words.", "speaker_uncertain": True}
        )
    elif failure == "provider":
        provider.chat.completions.create.side_effect = RuntimeError(
            "SYNTHETIC_PRIVATE_PROVIDER_ERROR"
        )
    else:
        error_type = RuntimeError if failure == "unexpected" else UntrustedTranscriptionError
        transcriber.enhance_transcription = MagicMock(
            side_effect=error_type("SYNTHETIC_PRIVATE_ERROR")
        )
    monkeypatch.setattr("src.cli.Transcriber", lambda **kwargs: transcriber)
    path = manifest(tmp_path, [synthetic_media("synthetic.wav")])
    assert (
        main(
            [
                "--workflow",
                "--interview-manifest",
                str(path),
                "--stages",
                "raw,polish",
                "--output-folder",
                str(tmp_path / "output"),
            ]
        )
        == 1
    )
    logs = capsys.readouterr()
    assert "SYNTHETIC_PRIVATE" not in logs.out + logs.err
    assert str(tmp_path) not in logs.out + logs.err
    message = json.loads(logs.out.splitlines()[-1])["message"]
    if failure == "fidelity":
        assert "words or symbols" in message
    elif failure == "provider":
        assert "API/model access, quota" in message
    else:
        assert message.startswith("Interview processing failed;")
