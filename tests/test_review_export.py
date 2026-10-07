"""Synthetic offline checks for lossless, private author-review workbooks."""

import copy
import hashlib
import json
import os
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from src.review_export import ReviewExportError, export_review


def _report(text='A synthetic account needs context.', *, status='complete'):
    report = {
        'status': status, 'raw_sha256': hashlib.sha256(text.encode()).hexdigest(),
        'model': 'synthetic-model', 'prompt_version': 'synthetic-v1', 'prompt_sha256': 'b' * 64,
        'schema_version': 1, 'settings_fingerprint': 'c' * 64,
        'segments': [{'segment_id': 'segment-0001', 'start': 0, 'end': len(text), 'text': text}],
        'findings': [{
            'finding_id': 'finding-0001', 'severity': 'medium', 'category': 'contextual_criticism',
            'excerpt': text, 'start': 0, 'end': len(text), 'segment_ids': ['segment-0001'],
            'rationale': 'A synthetic passage needs context.',
            'author_action': 'Verify the intended attribution with the author.',
            'author_question': 'What context should accompany this account?',
            'disposition': 'Unreviewed', 'reviewer': '',
        }],
        'coverage': {'complete': status == 'complete', 'total_characters': len(text),
                     'reviewed_characters': len(text) if status == 'complete' else 0,
                     'chunks': [{'chunk_index': 0, 'start': 0, 'end': len(text),
                                 'context_start': 0, 'context_end': len(text),
                                 'status': 'complete' if status == 'complete' else 'failed',
                                 'error': '' if status == 'complete' else 'provider_refusal'}]},
    }
    if status == 'incomplete':
        boundary = len(text) // 2
        report['coverage']['reviewed_characters'] = boundary
        report['coverage']['chunks'] = [
            {'chunk_index': 0, 'start': 0, 'end': boundary, 'context_start': 0,
             'context_end': len(text), 'status': 'complete'},
            {'chunk_index': 1, 'start': boundary, 'end': len(text), 'context_start': 0,
             'context_end': len(text), 'status': 'failed', 'error': 'provider_refusal'},
        ]
    return report


def _metadata(workbook):
    return dict(workbook['Overview'].iter_rows(min_row=2, values_only=True))


def test_report_roundtrip_preserves_text_refs_and_review_fields(tmp_path):
    report = _report('Synthetic ā̄ \u200d testimony.\nSecond line.')
    target = tmp_path / 'private' / 'review.xlsx'
    export_review(report, target)
    workbook = load_workbook(target)
    assert workbook.sheetnames == ['Overview', 'Findings', 'Source segments', 'Coverage']
    finding = list(workbook['Findings'].iter_rows(min_row=2, values_only=True))[0]
    assert finding == (
        'finding-0001', 'medium', 'contextual_criticism', report['findings'][0]['excerpt'],
        0, len(report['segments'][0]['text']), 'segment-0001',
        report['findings'][0]['rationale'], report['findings'][0]['author_question'],
        report['findings'][0]['author_action'], 'Unreviewed', None, 'Literal raw text',
    )
    assert _metadata(workbook)['Raw SHA-256'] == report['raw_sha256']
    assert workbook['Findings'].auto_filter.ref == 'A1:M2'
    assert workbook['Findings'].freeze_panes == 'A2'
    assert workbook['Findings']['B2'].fill.fgColor.rgb.endswith('FFF0C2')
    assert workbook['Findings']['K2'].fill.fgColor.rgb.endswith('E8F4FC')
    assert len(workbook['Findings'].data_validations.dataValidation) == 1
    assert target.stat().st_mode & 0o777 == 0o600
    assert target.parent.stat().st_mode & 0o777 == 0o700


@pytest.mark.parametrize('prefix', ['=', '+', '-', '@'])
def test_formula_like_source_and_model_cells_are_literal(tmp_path, prefix):
    report = _report(prefix + 'HYPERLINK("https://invalid.example","synthetic")')
    for key in ('rationale', 'author_action', 'author_question', 'reviewer'):
        report['findings'][0][key] = prefix + 'SYNTHETIC_FORMULA'
    report['model'] = prefix + 'SYNTHETIC_MODEL'
    target = tmp_path / 'report.xlsx'
    export_review(report, target)
    workbook = load_workbook(target, data_only=False)
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert cell.data_type != 'f'
                assert cell.hyperlink is None
    assert workbook['Findings']['D2'].value == report['findings'][0]['excerpt']
    with zipfile.ZipFile(target) as archive:
        for name in archive.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                assert b'<f>' not in archive.read(name)


@pytest.mark.parametrize('status', ['complete', 'incomplete', 'failed'])
def test_empty_complete_is_distinct_from_empty_failed_or_incomplete(tmp_path, status):
    report = _report(status=status)
    report['findings'] = []
    target = tmp_path / f'{status}.xlsx'
    export_review(report, target)
    metadata = _metadata(load_workbook(target))
    assert metadata['Report status'] == status
    assert metadata['Finding count'] == 0
    if status == 'complete':
        assert 'It does not establish accuracy' in metadata['Status meaning']
    else:
        assert 'not a clean review' in metadata['Status meaning']


def test_long_source_and_supplementary_unicode_split_losslessly(tmp_path):
    text = 'ā\u200d😀' * 16000
    report = _report(text)
    report['findings'] = []
    target = tmp_path / 'report.xlsx'
    export_review(report, target)
    rows = list(load_workbook(target)['Source segments'].iter_rows(min_row=2, values_only=True))
    assert len(rows) > 1
    assert ''.join(row[4] for row in rows) == text
    assert rows[0][2] == 0
    assert rows[-1][3] == len(text)
    for first, second in zip(rows, rows[1:]):
        assert first[3] == second[2]
    assert all(len(row[4].encode('utf-16-le')) // 2 <= 32767 for row in rows)


def test_xml_illegal_source_is_lossless_and_labelled(tmp_path):
    text = '=Synthetic\x00\x0b\uffff words.'
    report = _report(text)
    target = tmp_path / 'report.xlsx'
    export_review(report, target)
    workbook = load_workbook(target)
    assert workbook['Findings']['M2'].value == 'JSON string (escaped)'
    assert json.loads(workbook['Findings']['D2'].value) == text
    source_rows = list(workbook['Source segments'].iter_rows(min_row=2, values_only=True))
    assert source_rows[0][5] == 'JSON string (escaped)'
    assert json.loads(source_rows[0][4]) == text
    assert 'JSON report remains canonical' in _metadata(workbook)['Text encoding']


def test_long_controls_source_is_split_without_escaped_cell_overflow(tmp_path):
    text = '\x00' * 12000
    report = _report(text)
    report['findings'] = []
    target = tmp_path / 'report.xlsx'
    export_review(report, target)
    rows = list(load_workbook(target)['Source segments'].iter_rows(min_row=2, values_only=True))
    assert ''.join(json.loads(row[4]) for row in rows) == text
    assert all(len(row[4]) <= 32767 for row in rows)


@pytest.mark.parametrize('field', ['rationale', 'author_action', 'author_question', 'reviewer'])
def test_oversize_text_field_is_rejected_without_silent_truncation(tmp_path, field):
    report = _report()
    report['findings'][0][field] = '😀' * 17000
    target = tmp_path / 'report.xlsx'
    with pytest.raises(ReviewExportError, match='cell limit'):
        export_review(report, target)
    assert not target.exists()
    assert not list(tmp_path.glob('.review-*'))


@pytest.mark.parametrize('mutation', [
    lambda report: report.update(status='unknown'),
    lambda report: report['findings'][0].update(excerpt='Invented synthetic quotation.'),
    lambda report: report['findings'][0].update(start=1),
    lambda report: report['findings'][0].update(segment_ids=['missing']),
    lambda report: report['findings'][0].update(severity='critical'),
    lambda report: report['segments'][0].update(start=1),
    lambda report: report['coverage'].update(complete=False),
    lambda report: report['coverage'].update(reviewed_characters=1),
    lambda report: report['coverage'].update(chunks=[]),
    lambda report: report.update(raw_sha256='0' * 64),
    lambda report: report['findings'].append(copy.deepcopy(report['findings'][0])),
])
def test_invalid_report_or_hallucinated_source_is_rejected(tmp_path, mutation):
    report = _report()
    mutation(report)
    target = tmp_path / 'report.xlsx'
    with pytest.raises(ReviewExportError):
        export_review(report, target)
    assert not target.exists()


def test_cross_segment_excerpt_references_exact_covered_segments(tmp_path):
    report = _report('Synthetic first. Synthetic second.')
    text = report['segments'][0]['text']
    boundary = 17
    report['segments'] = [
        {'segment_id': 'segment-0001', 'start': 0, 'end': boundary, 'text': text[:boundary]},
        {'segment_id': 'segment-0002', 'start': boundary, 'end': len(text), 'text': text[boundary:]},
    ]
    report['findings'][0]['segment_ids'] = ['segment-0001', 'segment-0002']
    target = tmp_path / 'report.xlsx'
    export_review(report, target)
    assert load_workbook(target)['Findings']['G2'].value == 'segment-0001, segment-0002'


def test_no_overwrite_existing_file_or_symlink(tmp_path):
    target = tmp_path / 'report.xlsx'
    target.write_bytes(b'synthetic existing data')
    with pytest.raises(ReviewExportError, match='already exists'):
        export_review(_report(), target)
    assert target.read_bytes() == b'synthetic existing data'
    target.unlink()
    original = tmp_path / 'original.xlsx'
    original.write_bytes(b'synthetic existing data')
    target.symlink_to(original)
    with pytest.raises(ReviewExportError, match='already exists'):
        export_review(_report(), target)
    assert original.read_bytes() == b'synthetic existing data'


def test_failed_save_cleans_temp_and_sanitizes_private_details(monkeypatch, tmp_path):
    def fail(self, stream):
        stream.write(b'partial')
        raise OSError('SYNTHETIC_PRIVATE_TEXT /private/source/path credential')
    monkeypatch.setattr('src.review_export.Workbook.save', fail)
    with pytest.raises(ReviewExportError) as error:
        export_review(_report(), tmp_path / 'report.xlsx')
    assert 'SYNTHETIC_PRIVATE_TEXT' not in str(error.value)
    assert 'source/path' not in str(error.value)
    assert not (tmp_path / 'report.xlsx').exists()
    assert not list(tmp_path.glob('.review-*'))


def test_publication_race_does_not_overwrite(monkeypatch, tmp_path):
    target = tmp_path / 'report.xlsx'
    real_link = os.link
    def race(source, destination):
        Path(destination).write_bytes(b'synthetic competing output')
        real_link(source, destination)
    monkeypatch.setattr('src.review_export.os.link', race)
    with pytest.raises(ReviewExportError, match='already exists'):
        export_review(_report(), target)
    assert target.read_bytes() == b'synthetic competing output'
    assert not list(tmp_path.glob('.review-*'))


def test_repository_output_must_be_in_private_tree(tmp_path):
    (tmp_path / '.git').mkdir()
    target = tmp_path / 'public' / 'report.xlsx'
    with pytest.raises(ReviewExportError):
        export_review(_report(), target)
    assert not target.exists()
    export_review(_report(), tmp_path / 'private' / 'report.xlsx')


def test_symlink_parent_and_wrong_extension_rejected(tmp_path):
    directory = tmp_path / 'actual'
    directory.mkdir()
    link = tmp_path / 'linked'
    link.symlink_to(directory, target_is_directory=True)
    with pytest.raises(ReviewExportError):
        export_review(_report(), link / 'report.xlsx')
    with pytest.raises(ReviewExportError, match='extension'):
        export_review(_report(), directory / 'report.csv')
    assert not list(directory.iterdir())


def test_actual_review_report_contract_exports_with_mocked_model(tmp_path):
    from src.author_review import ReviewOptions, review_transcript

    client = MagicMock()
    def create(**kwargs):
        payload = json.loads(kwargs['messages'][-1]['content'])
        content = json.dumps({
            'chunk_index': payload['chunk_index'], 'fully_reviewed': True,
            'contract_version': payload['contract_version'], 'reviewed_piece_ids': payload['core_piece_ids'],
            'findings': [],
        })
        return SimpleNamespace(choices=[SimpleNamespace(
            finish_reason='stop', message=SimpleNamespace(content=content, refusal=None))])
    client.chat.completions.create.side_effect = create
    raw = 'Synthetic first line.\r\nSynthetic second line ā\u200d😀.\n'
    report = review_transcript(raw, client, ReviewOptions(chunk_bytes=96))
    target = tmp_path / 'report.xlsx'
    export_review(report, target)
    workbook = load_workbook(target)
    metadata = _metadata(workbook)
    assert metadata['Report status'] == 'complete'
    assert metadata['Schema version'] == report['schema_version']
    assert metadata['Settings fingerprint'] == report['settings_fingerprint']
    assert ''.join(json.loads(row[4]) if row[5] == 'JSON string (escaped)' else row[4]
                   for row in workbook['Source segments'].iter_rows(
                       min_row=2, values_only=True)) == raw


def _recording_report(raw, client, *, part_id='synthetic-part'):
    from src.author_review import review_transcript, source_segments
    from src.source_provenance import bind_report
    report = review_transcript(raw, client)
    checksum = hashlib.sha256(raw.encode()).hexdigest()
    provenance = {
        'version': 1, 'human_review_required': True, 'separator': '\n\n',
        'raw_sha256': checksum,
        'parts': [{'id': part_id, 'order': 1, 'start': 0, 'end': len(raw),
                   'path': '=SYNTHETIC_SOURCE.wav', 'source_sha256': 'a' * 64,
                   'raw_transcript': '=SYNTHETIC_TRANSCRIPT', 'raw_sha256': checksum,
                   'segments': [{k: v for k, v in segment.items() if k != 'text'}
                                for segment in source_segments(raw)]}],
        'spans': [{'kind': 'recording', 'part_id': part_id, 'order': 1, 'start': 0, 'end': len(raw)}],
    }
    bind_report(report, provenance)
    return report


def test_recording_citation_cell_limit_rejects_incomplete_export(tmp_path):
    client = MagicMock()
    def respond(**kwargs):
        data = json.loads(kwargs['messages'][-1]['content'])
        body = {'chunk_index': data['chunk_index'], 'fully_reviewed': True, 'findings': [],
                'contract_version': data['contract_version'], 'reviewed_piece_ids': data['core_piece_ids']}
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',
            message=SimpleNamespace(content=json.dumps(body), refusal=None))])
    client.chat.completions.create.side_effect = respond
    report = _recording_report('a\n' * 3000, client)
    assert len(', '.join(s['segment_id'] for s in report['recording_provenance']['parts'][0]['segments'])) > 32767
    target = tmp_path / 'synthetic-overflow.xlsx'
    with pytest.raises(ReviewExportError, match='cell limit'):
        export_review(report, target)
    assert not target.exists()


def test_recording_citations_below_cell_limit_and_literal_formula_prefixes(tmp_path):
    client = MagicMock()
    def respond(**kwargs):
        data = json.loads(kwargs['messages'][-1]['content'])
        body = {'chunk_index': data['chunk_index'], 'fully_reviewed': True, 'findings': [],
                'contract_version': data['contract_version'], 'reviewed_piece_ids': data['core_piece_ids']}
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',
            message=SimpleNamespace(content=json.dumps(body), refusal=None))])
    client.chat.completions.create.side_effect = respond
    report = _recording_report('a\n' * 2730, client, part_id='=SYNTHETIC_PART')
    expected = ', '.join(s['segment_id'] for s in report['recording_provenance']['parts'][0]['segments'])
    assert len(expected) == 32758
    target = tmp_path / 'synthetic-safe-citations.xlsx'
    export_review(report, target)
    workbook = load_workbook(target)
    chunk = next(row for row in workbook['Recording references'].iter_rows(min_row=2) if row[0].value == 'chunk')
    assert chunk[9].value == expected
    assert chunk[4].value == '=SYNTHETIC_PART'
    assert workbook['Recording parts']['C2'].value == '=SYNTHETIC_SOURCE.wav'
    for name in ('Recording parts', 'Recording references'):
        assert all(cell.data_type != 'f' for row in workbook[name] for cell in row)
